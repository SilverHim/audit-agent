"""需求 R338 (Issue #50): Xero 演示的工作底稿接入.

演示配置走通 (tests/test_demo_walkthrough.walk) 到归档后:
- 档案接口按索引列出全部底稿, 已生成的都定稿为 final 且带未签署标记; 整套 zip 的条目与索引同名, 条目数等于生成份数,
  文件显式 UTF-8, 常见解压工具里中文文件名不乱码, 解压出的 Excel 可读.
- D-3, D-2, A-2, D, X, S 六份与 docs/云端开发/design/底稿示例/ 的示例逐份比对: 表头各项, 节的顺序与表的栏目按示例;
  数据按本次走通的实际结果 (A-2 的执行重要性按演示决定 50% 而不是示例的 75%).
- 每个执行过的程序在完成后都有底稿 (test_workpapers 的版本历史用例覆盖), 本节另核索引里的状态与功能未开启项.
"""

import re
from collections.abc import AsyncIterator
from io import BytesIO
from pathlib import Path
from zipfile import ZipFile

import httpx
import pytest
import pytest_asyncio
from openpyxl import load_workbook

from backend.contracts.models import Workpaper, WorkpaperStatus
from tests.fixture_sets import XERO
from tests.test_demo_walkthrough import Session, Walk, walk
from tests.xero_flow import demo_offline
from tests.xero_harness import LayeredMockLLMClient

SAMPLE_DIR = (
    Path(__file__).resolve().parents[1] / "docs" / "云端开发" / "design" / "底稿示例" / "ENG-2026-001 审计档案 (示例)"
)
SAMPLES = {
    "00": "00 索引.xlsx",
    "A-2": "A-2 重要性 (PLN-01).xlsx",
    "D": "D 收入导引表.xlsx",
    "D-2": "D-2 收入合同条款评价 (REV-03).xlsx",
    "D-3": "D-3 凭证三单匹配 (REV-08).xlsx",
    "X": "X 错报汇总.xlsx",
    "S": "S 准则对照.xlsx",
}
# 示例的节 (小节标题) 与生成内容的节 id 对应: 目标与风险 / 依据 / 程序 / 检查项 / 测试明细 / 例外与处理 / 结论 /
# 交叉索引 / 复核意见. 示例把 "依据" 并入 目标与风险 的 D-3, D-2 两份在生成内容里单列一节, 顺序不变.
SECTIONS = {
    "00": ["details", "conclusion", "review"],
    "A-2": [
        "objectives",
        "basis",
        "procedures",
        "checks",
        "details",
        "exceptions",
        "conclusion",
        "cross_references",
        "review",
    ],
    "D": [
        "objectives",
        "basis",
        "procedures",
        "checks",
        "details",
        "exceptions",
        "conclusion",
        "cross_references",
        "review",
    ],
    "D-2": [
        "objectives",
        "basis",
        "procedures",
        "checks",
        "details",
        "exceptions",
        "conclusion",
        "cross_references",
        "review",
    ],
    "D-3": [
        "objectives",
        "basis",
        "procedures",
        "checks",
        "details",
        "exceptions",
        "conclusion",
        "cross_references",
        "review",
    ],
    "X": [
        "objectives",
        "basis",
        "procedures",
        "checks",
        "details",
        "exceptions",
        "conclusion",
        "cross_references",
        "review",
    ],
    "S": ["details", "conclusion", "review"],
}
# 各份的关键表的栏目 (中文标签): 与示例的栏目对应, 名称差异只限 程序/对应程序 这类措辞 (R338 验收 2 的结构比对).
COLUMNS = {
    "00": {
        "workpaper_index": [
            "索引",
            "底稿",
            "程序",
            "编制人",
            "完成日期",
            "状态",
            "复核人",
            "复核时间",
            "复核范围",
            "缺项数",
        ]
    },
    "A-2": {
        "materiality_candidates": ["比较基准", "基准金额", "比例下限", "比例上限", "金额下限", "金额上限", "状态"],
        "materiality_calculation": [
            "编号",
            "比较基准",
            "基准金额",
            "整体重要性比例",
            "整体重要性",
            "执行比例",
            "执行重要性",
            "明显微小比例",
            "明显微小临界值",
            "特定类别重要性",
            "理由",
        ],
    },
    "D": {
        "lead_schedule": [
            "科目",
            "科目名称",
            "本期未审数",
            "已更正调整",
            "审定数",
            "上期末余额",
            "变动",
            "未更正错报",
            "来源底稿",
        ]
    },
    "D-2": {"judgment_summary": ["对象", "编号", "结论", "置信度", "建议", "状态"]},
    "D-3": {
        "tested_items": [
            "发票号",
            "客户",
            "合同",
            "入账日期",
            "金额",
            "分录",
            "服务完成记录",
            "签收日",
            "签收人",
            "收款流水",
            "收款日",
            "收款金额",
            "入账期间正确",
            "有服务完成记录",
            "收款金额相符",
            "备注",
        ],
        "exceptions": ["标记", "编号", "检查项", "对象", "差额", "状态", "判断"],
    },
    "X": {
        "misstatements": [
            "索引",
            "编号",
            "来源底稿",
            "对象",
            "科目",
            "认定",
            "类型",
            "方向",
            "金额",
            "超过明显微小",
            "状态",
            "累计未更正金额",
            "决定人",
            "决定时间",
        ],
        "misstatement_totals": ["合计", "执行重要性", "整体重要性", "占执行重要性", "占整体重要性", "明显微小临界值"],
    },
    "S": {"standards": ["准则", "段号", "说明", "来源底稿", "状态", "备注"]},
}


def sample_rows(ref: str) -> list[list[object]]:
    """读示例工作簿的有效行 (非空单元格), 供与生成内容按数据比对."""
    workbook = load_workbook(SAMPLE_DIR / SAMPLES[ref], read_only=True)
    sheet = workbook.active
    rows = [[cell for cell in row if cell is not None] for row in sheet.iter_rows(values_only=True)]
    workbook.close()
    return [row for row in rows if row]


def sample_item(row: list[object]) -> tuple[str, str, str | None, str, int]:
    """示例 D-3 测试明细的一行: (发票号, 客户, 合同, 入账日期, 金额分). 合同为 — 时记 None."""
    contract = row[3]
    return (
        str(row[1]),
        str(row[2]),
        None if contract == "—" else str(contract),
        str(row[4]),
        round(float(row[5]) * 100),
    )


def sample_numbers(sheet: list[list[object]], label: str) -> float:
    """示例里首列为 label 的那一行, 第二个单元格的数值 (如 整体重要性 324.31)."""
    [row] = [row for row in sheet if row[0] == label]
    value = row[1]
    assert isinstance(value, int | float)
    return float(value)


def tables_of(paper: Workpaper) -> dict[str, object]:
    """生成内容里的表, 表 id → WorkpaperTable."""
    assert paper.content is not None
    return {table.id: table for section in paper.content.sections for table in section.tables}


def zh(value: object) -> str:
    """LocalizedText 或普通值取中文标签."""
    return value.zh if hasattr(value, "zh") else str(value)


@pytest_asyncio.fixture(scope="module", loop_scope="module")
async def walked() -> AsyncIterator[tuple[Walk, httpx.AsyncClient]]:
    """演示配置离线走通一次 (同 tests/test_demo_walkthrough 的 walked, 工作语言 en), 走完不关接口客户端."""
    XERO.ready()
    llm = LayeredMockLLMClient([XERO.runtime])
    async with demo_offline(llm) as (runner, service, client):
        yield await walk(Session(client, runner, service), llm), client


shared_loop = pytest.mark.asyncio(loop_scope="module")


@pytest.mark.external
@shared_loop
async def test_archive_lists_and_zips_every_generated_workpaper(walked: tuple[Walk, httpx.AsyncClient]) -> None:
    """验收 1: 档案接口列出全部底稿, 已生成的都是 final 且未签署; 整套 zip 的条目与索引同名, 条目数等于生成份数,
    显式 UTF-8 且解压出的 Excel 可读."""
    w, client = walked
    base = f"/api/v1/engagements/{w.x.eid}/workpapers"
    listed = (await client.get(base)).json()
    refs = w.x.runner.deps.library.workpaper_refs()
    assert [row["ref"] for row in listed] == refs
    generated = [row for row in listed if row["availability"] == "generated"]
    disabled = {row["ref"] for row in listed if row["availability"] == "disabled"}
    assert disabled == {"B-3", "G-3", "G-4", "G-5", "G-6"}  # 演示配置未开 acceptance, related_parties 与 going_concern
    assert {row["status"] for row in generated} == {"final"}
    assert all(row["unsigned"] is True for row in generated)  # 本走通不签署底稿, 未签署仍定稿
    assert all(row["file_name"] for row in generated)

    response = await client.get(f"{base}/archive")
    assert response.status_code == 200 and response.headers["content-type"] == "application/zip"
    bundle = ZipFile(BytesIO(response.content))
    names = [info.filename for info in bundle.infolist()]
    assert names == [row["file_name"] for row in generated]
    # 中文文件名按 UTF-8 显式标记
    assert all(info.flag_bits & 0x800 for info in bundle.infolist())
    d3_name = next(n for n in names if n.startswith("D-3 "))
    d3 = load_workbook(BytesIO(bundle.read(d3_name)), read_only=True)
    assert str(d3.active["A1"].value).startswith("D-3  ")  # 标题随工作语言 (走通为 en) 渲染, 索引号与节名同索引
    d3.close()
    single = await client.get(f"{base}/D-3/file")
    assert single.status_code == 200 and single.content[:2] == b"PK"


@pytest.mark.external
@shared_loop
async def test_structure_follows_the_confirmed_samples(walked: tuple[Walk, httpx.AsyncClient]) -> None:
    """验收 2 (结构): 六份与索引的表头各项, 节的顺序, 关键表的栏目与示例一致; 标题与索引号取示例."""
    w, _ = walked
    papers = {p.ref: p for p in await w.x.runner.deps.output_store.latest(w.x.eid, Workpaper)}
    for ref, sections in SECTIONS.items():
        paper = papers[ref]
        assert paper.status is WorkpaperStatus.final and paper.content is not None
        header = paper.content.header
        assert header.title.zh and header.title.en
        assert (header.client_name, header.engagement_id) == ("Demo Company (AU) Pty Ltd", "ENG-2026-001")
        assert (header.period_start.isoformat(), header.period_end.isoformat()) == ("2026-07-01", "2026-09-30")
        assert header.currency == "AUD"
        assert [section.id for section in paper.content.sections] == sections
        for table_id, labels in COLUMNS.get(ref, {}).items():
            table = tables_of(paper)[table_id]
            assert [zh(column.label) for column in table.columns] == labels
    # 示例的表头 (客户, 审计期间, 项目, 对应程序, 索引) 与生成内容一致
    assert papers["D-3"].content.header.procedure_id == "REV-08"
    assert papers["D-2"].content.header.procedure_id == "REV-03"
    assert papers["A-2"].content.header.procedure_id == "PLN-01"
    assert [papers[ref].content.header.procedure_id for ref in ("X", "S", "00")] == ["FIN-04", None, None]
    assert sample_rows("D-3")[1][0].startswith("客户: Demo Company (AU) Pty Ltd")


@pytest.mark.external
@shared_loop
async def test_data_matches_the_samples_and_this_run(walked: tuple[Walk, httpx.AsyncClient]) -> None:
    """验收 2 (数据): D-3 的 6 张发票与三项检查, D-2 的 3 份合同, A-2 的重要性, D 的 3 个科目, X 的 2 条错报,
    S 的 15 行准则与示例一致 (段号按原文, ASA 240 为 45-48; 状态按本次走通)."""
    w, _ = walked
    papers = {p.ref: p for p in await w.x.runner.deps.output_store.latest(w.x.eid, Workpaper)}

    # D-3: 测试明细的 6 张发票 (示例从 23 行起, 列为 序号, 发票号, 客户, 合同, 入账日期, 金额), 三项检查结果与例外
    sheet = sample_rows("D-3")
    expected = [sample_item(row) for row in sheet if isinstance(row[0], int)]
    table = tables_of(papers["D-3"])["tested_items"]
    actual = [
        (row["invoice_no"], row["customer"], row["contract_id"], row["posting_date"], row["amount_cents"])
        for row in table.rows
    ]
    assert expected == actual
    assert [row["check_RUN-REV-08-1_R01"] for row in table.rows] == ["✓"] * 6
    assert [row["check_RUN-REV-08-1_R02"] for row in table.rows] == ["✓"] * 6
    receipts = [row["check_RUN-REV-08-1_R03"] for row in table.rows]
    assert (receipts.count("✓"), receipts.count("✗"), receipts.count("—")) == (4, 1, 1)  # 5 项中 1 不符, 1 项不作判断
    [exception] = tables_of(papers["D-3"])["exceptions"].rows
    assert (exception["id"], exception["difference_cents"], exception["marker"]) == ("X-REV-08-1-R03-1", 330000, "a")

    # D-2: 3 份合同 (示例的 3.x 小节标题), 建议接受, 置信度同示例的 0.84
    contracts = [
        match.group(1) for row in sample_rows("D-2") if (match := re.match(r"^3\.\d+\s+(CON-X\d{2})", str(row[0])))
    ]
    assert contracts == ["CON-X01", "CON-X02", "CON-X03"]
    summary = tables_of(papers["D-2"])["judgment_summary"].rows
    assert [row["source_id"] for row in summary] == ["CON-X01", "CON-X02", "CON-X03"]
    assert {zh(row["recommendation"]) for row in summary} == {"接受"}
    assert [row["confidence"] for row in summary] == [0.84, 0.84, 0.82]

    # A-2: 整体重要性与明显微小临界值同示例; 执行重要性按演示决定 (50%), 示例取 75%
    sheet = sample_rows("A-2")
    calculation = tables_of(papers["A-2"])["materiality_calculation"].rows[0]
    assert calculation["base_amount_cents"] == int(sample_numbers(sheet, "基准: 收入") * 100)
    overall = round(sample_numbers(sheet, "整体重要性") * 100)
    trivial = round(sample_numbers(sheet, "明显微小临界值") * 100)
    assert calculation["overall_cents"] == overall
    assert calculation["trivial_cents"] == trivial
    assert calculation["performance_cents"] == 16216  # 演示决定 50% (示例 75% 是 243.23 的示例值)
    assert zh(calculation["basis"]) == "收入" and calculation["pct"] == 0.01

    # D: 示例的 3 个科目 200 / 260 / 270 与名称
    sheet = sample_rows("D")
    expected = [(str(row[0]), str(row[1])) for row in sheet if re.fullmatch(r"\d{3}", str(row[0]))]
    lead = tables_of(papers["D"])["lead_schedule"].rows
    assert [(row["account_code"], row["account_name"]) for row in lead] == expected

    # X: 本走通认定两条错报 (BA-091 48.53 低估, XJ-499 750.00 高估), 都已更正, 汇总 0 且与候选意见一致
    misstatements = tables_of(papers["X"])["misstatements"].rows
    assert {(row["amount_cents"], zh(row["direction"]), zh(row["status"])) for row in misstatements} == {
        (4853, "低估", "已更正"),
        (75000, "高估", "已更正"),
    }
    totals = tables_of(papers["X"])["misstatement_totals"].rows[0]
    assert totals["total_cents"] == 0 and totals["overall_cents"] == 32431 and totals["performance_cents"] == 16216

    # S: 示例的 15 行准则与段落, 按仓库原文核对段号 (ASA 240 示例写 "45 起", 这里写 45-48)
    sheet = [row for row in sample_rows("S") if re.fullmatch(r"ASA \d+", str(row[0]))]
    assert len(sheet) == 15
    paragraphs = {str(row[0]): str(row[1]) for row in sheet}
    paragraphs["ASA 240"] = "45-48"
    standards = tables_of(papers["S"])["standards"].rows
    assert [(row["standard"], row["paragraph"]) for row in standards] == list(paragraphs.items())
    assert all(row["standard"] and row["paragraph"] and zh(row["status"]) for row in standards)


@pytest.mark.external
@shared_loop
async def test_index_lists_every_workpaper_in_file_order(walked: tuple[Walk, httpx.AsyncClient]) -> None:
    """验收 1 的索引部分: 00 索引的 32 行按档案顺序, 每行有索引号与底稿名; 已生成的行标最终版, 未开启功能的标未开启."""
    w, _ = walked
    papers = {p.ref: p for p in await w.x.runner.deps.output_store.latest(w.x.eid, Workpaper)}
    index = tables_of(papers["00"])["workpaper_index"].rows
    refs = w.x.runner.deps.library.workpaper_refs()
    assert [row["ref"] for row in index] == refs
    disabled = {row["ref"] for row in index if zh(row["status"]) == "功能未开启"}
    assert disabled == {"B-3", "G-3", "G-4", "G-5", "G-6"}
    assert all(row["title"].zh and row["title"].en for row in index)
