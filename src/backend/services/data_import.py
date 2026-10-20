"""账表导入的解析 (需求 R13): 读 CSV, XLSX, XLS, 建议表种类与列对应, 各账表的导入字段与单元格的类型转换.
按列对应或规整方案把各行转成财务表的行模型在 services/table_plan.

只做解析与类型转换, 不做业务校验 (借贷平衡, 与试算表一致等由 A3 体检负责). 规则:
- 原样读行 (raw_rows) 保留标题行与空行; 表头缺省取第一个非空行 (规整方案可指定表头行), 之后的非空行为数据行;
  行号是文件中的实际行号 (第 1 行为 1).
- 金额一律存整数分: 源列的原始表头以 cents 结尾 (忽略大小写与符号) 时按整数分读, 否则按元读, 可带 $, AUD,
  千分位, 括号负数, 末尾负号, 四舍五入到分; 超出 BIGINT 算失败. Excel 的数值单元格同样按此处理.
- 日期: Excel 日期单元格直接取; 文本年在前的总按年月日, 三段数字按 date_format, 也认 "1 Jul 2025" 这类月份名;
  整个单元格须是这样的写法 (可带合法的时间与时区), 多出的内容算失败.
- 日期时间: 带时区的按其时区, 不带的按澳洲悉尼时间 (夏令时跳过的时刻算失败), 存 UTC; 只有日期时取当天零点.
- 编号列没对应时按确定的规则生成 (见 GENERATED_IDS), 同一文件重导结果相同.
"""

import csv
import io
import re
from collections.abc import Callable, Iterable, Iterator, Mapping, Sequence
from dataclasses import dataclass
from datetime import UTC, date, datetime, time, timedelta, timezone
from decimal import ROUND_HALF_UP, Decimal, DecimalException, InvalidOperation
from pathlib import PurePath
from typing import Any, cast
from zoneinfo import ZoneInfo

from pydantic import BaseModel

from backend.contracts.data import (
    ApprovalLogRow,
    DateFormat,
    FieldType,
    ImportPreview,
    TableField,
    TableKind,
    TableSchema,
)
from backend.contracts.records import (
    Account,
    BankAccount,
    BankTransaction,
    DeliveryNote,
    Employee,
    JournalEntry,
    SalesInvoice,
    SupplierBill,
    TrialBalanceRow,
)

MAX_BYTES = 50 * 1024 * 1024  # 单个文件上限 50 MB
EXTENSIONS = {".csv", ".xlsx", ".xlsm", ".xls"}  # 支持的文件
MAX_FAILURES = 20  # IMPORT_PARSE_ERROR 附带的失败处数
MAX_CENTS = 2**63 - 1  # 金额列为 BIGINT, 超出即解析失败
SYDNEY = ZoneInfo("Australia/Sydney")  # 不带时区的日期时间按悉尼时间

Cell = str | int | float | bool | date | datetime | None  # 单元格的值: CSV 都是文本, Excel 保留数值与日期


def _fields(spec: str) -> list[TableField]:
    """ "name:type!" 列表 → 字段, ! 表示必须对应. 类型缩写: t 文本, d 日期, dt 日期时间, a 金额, b 是否, i 整数."""
    types: dict[str, FieldType] = {
        "t": "text",
        "d": "date",
        "dt": "datetime",
        "a": "amount",
        "b": "bool",
        "i": "integer",
    }
    result = []
    for item in spec.split():
        name, kind = item.split(":")
        result.append(TableField(name=name, required=kind.endswith("!"), type=types[kind.rstrip("!")]))
    return result


# 各账表可对应的字段, 名称与 persistence/models.py 中 FINANCIAL 的列一致 (科目分类由系统归类, 不导入).
# 编号列 (journal_entries, sales_invoices, supplier_bills, delivery_notes, bank_transactions 的 id) 可不对应, 按
# GENERATED_IDS 生成
SCHEMAS: dict[TableKind, TableSchema] = {
    kind: TableSchema(kind=kind, fields=_fields(spec))
    for kind, spec in {
        TableKind.accounts: "code:t! name:t! type:t!",
        TableKind.trial_balance: "period:t! account_code:t! balance_cents:a!",
        TableKind.journal_entries: (
            "id:t entry_no:t! posting_date:d! created_at:dt created_by:t approved_by:t approved_at:dt "
            "account_code:t! debit_cents:a credit_cents:a description:t source_doc:t is_adjustment:b"
        ),
        TableKind.sales_invoices: (
            "id:t invoice_no:t! customer:t! customer_abn:t contract_id:t invoice_date:d! posting_date:d! "
            "amount_cents:a! journal_entry_id:t"
        ),
        TableKind.supplier_bills: (
            "id:t bill_no:t! supplier:t! bill_date:d! due_date:d! amount_cents:a! outstanding_cents:a! "
            "journal_entry_id:t"
        ),
        TableKind.bank_accounts: "id:t! bank_name:t! bsb:t! account_no:t! gl_account_code:t!",
        TableKind.bank_transactions: (
            "id:t bank_account_id:t! txn_date:d! amount_cents:a! balance_after_cents:a! description:t reference:t"
        ),
        TableKind.delivery_notes: "id:t delivery_no:t! invoice_id:t! shipped_date:d signed_date:d signed_by:t",
        TableKind.approval_log: "entry_no:t! created_by:t approved_by:t approved_at:dt",
        TableKind.employees: (
            "employee_id:t! name:t! start_date:d! employment_type:t pay_period_days:i! gross_per_period_cents:a!"
        ),
    }.items()
}
# 至少要对应其中一个的字段组: 分录的借贷, 审批日志补进总账的三列
ONE_OF: dict[TableKind, tuple[str, ...]] = {
    TableKind.journal_entries: ("debit_cents", "credit_cents"),
    TableKind.approval_log: ("created_by", "approved_by", "approved_at"),
}
ROW_MODELS: dict[TableKind, type[BaseModel]] = {
    TableKind.accounts: Account,
    TableKind.trial_balance: TrialBalanceRow,
    TableKind.journal_entries: JournalEntry,
    TableKind.sales_invoices: SalesInvoice,
    TableKind.supplier_bills: SupplierBill,
    TableKind.bank_accounts: BankAccount,
    TableKind.bank_transactions: BankTransaction,
    TableKind.delivery_notes: DeliveryNote,
    TableKind.approval_log: ApprovalLogRow,
    TableKind.employees: Employee,
}
# 唯一键: 同一文件里重复即解析失败 (写不进表); 审批日志一张凭证一行
KEYS: dict[TableKind, tuple[str, ...]] = {
    TableKind.accounts: ("code",),
    TableKind.trial_balance: ("period", "account_code"),
    TableKind.journal_entries: ("id",),
    TableKind.sales_invoices: ("id",),
    TableKind.supplier_bills: ("id",),
    TableKind.bank_accounts: ("id",),
    TableKind.bank_transactions: ("id",),
    TableKind.delivery_notes: ("id",),
    TableKind.approval_log: ("entry_no",),
    TableKind.employees: ("employee_id",),
}

# 列名的别名 (规范化后: 小写, 只留字母数字), 按优先顺序; 字段名本身 (含去掉 cents 的写法) 总是最先匹配
ALIASES: dict[str, tuple[str, ...]] = {
    "code": ("accountcode", "accountno", "accountnumber", "acctcode", "number"),
    "name": ("accountname", "title", "account"),
    "type": ("accounttype", "class", "accountclass"),
    "period": ("year", "fy"),
    "account_code": ("accountcode", "account", "code", "acctcode", "accountno", "glcode", "glaccount"),
    "balance_cents": ("balance", "closingbalance", "net", "ytdbalance", "amount"),
    "id": ("lineid", "rowid"),
    "entry_no": ("entryno", "journalno", "journalnumber", "entrynumber", "voucherno", "voucher"),
    "posting_date": ("postingdate", "postdate", "gldate", "journaldate", "date"),
    "created_at": ("createdat", "createddate", "created", "enteredat", "timestamp"),
    "created_by": ("createdby", "preparedby", "preparer", "enteredby", "user"),
    "approved_by": ("approvedby", "approver", "authorisedby", "authorizedby"),
    "approved_at": ("approvedat", "approvaldate", "approveddate", "authorisedat", "authorizedat"),
    "debit_cents": ("debit", "dr", "debitamount"),
    "credit_cents": ("credit", "cr", "creditamount"),
    "description": ("narration", "memo", "details", "particulars"),
    "source_doc": ("sourcedocument", "documentno", "docno", "reference", "ref"),
    "is_adjustment": ("adjustment", "adjusting", "adj"),
    "invoice_no": ("invoicenumber", "invoice", "number"),
    "customer": ("customername", "contact", "contactname", "client", "debtor"),
    "customer_abn": ("abn", "contactabn", "taxnumber"),
    "contract_id": ("contract", "contractno"),
    "invoice_date": ("issuedate", "date"),
    "amount_cents": ("amount", "total", "invoicetotal", "totalamount", "gross", "value"),
    "journal_entry_id": ("journalid", "journalline"),
    "bill_no": ("billnumber", "billno"),
    "supplier": ("suppliername", "vendor", "vendorname", "creditor", "payee"),
    "bill_date": ("billdate",),
    "due_date": ("duedate", "paymentdue", "dueby"),
    "outstanding_cents": ("owing", "unpaid"),  # "Outstanding" 由字段名本身对上
    "delivery_no": ("deliverynumber", "dispatchno", "shipmentno", "dnno"),
    "invoice_id": ("invoice", "invoiceno", "invoicenumber"),
    "shipped_date": ("shipdate", "dispatchdate", "dispatched"),
    "signed_date": ("receiveddate", "deliverydate"),
    "signed_by": ("receivedby",),
    "bank_name": ("bank", "institution"),
    "bsb": ("bsbnumber", "branch"),
    "account_no": ("accountnumber", "account"),
    "gl_account_code": ("glaccount", "glcode", "ledgeraccount"),
    "bank_account_id": ("bankaccount", "accountid", "account"),
    "txn_date": ("transactiondate", "valuedate", "date"),
    "balance_after_cents": ("balanceafter", "runningbalance", "balance", "closingbalance"),
    "reference": ("ref",),
    "employee_id": ("employeeid", "employeeno", "employeenumber", "staffid", "staffno", "payrollid"),
    "start_date": ("startdate", "commencementdate", "hiredate", "datejoined"),
    "employment_type": ("employmenttype", "employmentbasis", "basis"),
    "pay_period_days": ("payperioddays", "payperiod", "payfrequencydays", "periodlength"),
    "gross_per_period_cents": ("grossperperiod", "grosspay", "grosswages", "gross", "ordinarypay"),
}
# 应付账单与应收清单共有的列名 (Xero 的导出两边都写 Contact, Invoice Number, Invoice Date, Amount Due): 选定应付账单后
# 照样按它们对应列, 但猜表种类时只在文件名有应付账单的提示 (NAME_HINTS) 时才算应付账单的证据, 免得常见的应收清单
# 被建议成应付账单 (需求 R178 审核 02)
SHARED_ALIASES: dict[TableKind, dict[str, tuple[str, ...]]] = {
    TableKind.supplier_bills: {
        "bill_no": ("invoicenumber", "invoiceno", "number", "reference"),
        "supplier": ("contact", "contactname"),
        "bill_date": ("invoicedate", "issuedate", "date"),
        "outstanding_cents": ("amountdue", "balancedue", "due"),
    },
}
# 文件名提示: 规范化后的文件名含这些词时加分
NAME_HINTS: dict[TableKind, tuple[str, ...]] = {
    TableKind.accounts: ("chartofaccounts", "accounts", "coa"),
    TableKind.trial_balance: ("trialbalance", "tb"),
    TableKind.journal_entries: ("journal", "generalledger", "gl"),
    TableKind.sales_invoices: ("invoice", "sales"),
    TableKind.supplier_bills: ("supplierbills", "bills", "payable", "accountspayable", "agedpayables"),
    TableKind.bank_accounts: ("bankaccounts",),
    TableKind.bank_transactions: ("banktransactions", "bankstatement", "transactions", "statement"),
    TableKind.delivery_notes: ("delivery", "dispatch", "shipment"),
    TableKind.approval_log: ("approval",),
    TableKind.employees: (),  # 必须字段别的表都对不全, 不会与别的表平分, 不需要文件名提示
}


# 扩展名 → 媒体类型, 原件登记为证据文件时用; 认不出的为 application/octet-stream
MEDIA_TYPES = {
    ".pdf": "application/pdf",
    ".csv": "text/csv",
    ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    ".xlsm": "application/vnd.ms-excel.sheet.macroEnabled.12",
    ".xls": "application/vnd.ms-excel",
    ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    ".doc": "application/msword",
    ".png": "image/png",
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
}


def media_type(file_name: str) -> str:
    """按扩展名 (忽略大小写) 取媒体类型."""
    return MEDIA_TYPES.get(PurePath(file_name).suffix.lower(), "application/octet-stream")


def normalize(name: str) -> str:
    """列名规范化: 小写, 只留字母与数字."""
    return re.sub(r"[^0-9a-z]", "", name.lower())


# ── 读文件 ──


@dataclass(frozen=True)
class Sheet:
    """一个工作表: 表头与数据行 (行号, 各列的值, 与表头对齐; 短行补空)."""

    name: str  # 工作表名
    header_row: int  # 表头所在行号; 工作表全空时为 0
    columns: list[str]  # 各列的显示名, 全表唯一 (同名的依次加 " (2)", " (3)" 等, 避开已有的列名)
    headers: list[str]  # 各列的原始表头文本 (空表头为空串), 金额按分还是按元读看它, 不看显示名
    rows: list[tuple[int, list[Cell]]]  # (文件中的行号, 各列的值)


class UnreadableFile(ValueError):
    """文件格式不支持或读不了, 接口报 INVALID_VALUE."""


class MappingError(ValueError):
    """列对应或规整方案与工作表对不上 (表头行, 字段, 列, 必须字段); code 为接口错误码 (INVALID_VALUE 或
    MISSING_FIELD)."""

    def __init__(self, code: str, message: str) -> None:
        super().__init__(message)
        self.code = code
        self.message = message


def extension(file_name: str) -> str:
    """文件扩展名 (小写, 带点); 不支持的抛 UnreadableFile."""
    ext = PurePath(file_name).suffix.lower()
    if ext not in EXTENSIONS:
        raise UnreadableFile(f"只支持 CSV, XLSX, XLS 文件: {file_name}")
    return ext


def sheet_names(file_name: str, data: bytes) -> list[str]:
    """全部工作表名; CSV 只有一个, 名为文件名主干."""
    ext = extension(file_name)
    if ext == ".csv":
        return [PurePath(file_name).stem]
    if ext == ".xls":
        return cast(list[str], _xls_book(data).sheet_names())
    workbook = _xlsx_book(data)
    try:
        return list(workbook.sheetnames)
    finally:
        workbook.close()


RawRows = list[tuple[int, list[Cell]]]  # 原样读出的行: (文件中的行号, 各单元格的值), 含标题行与空行


def raw_rows(file_name: str, data: bytes, sheet: str | None = None) -> tuple[str, RawRows]:
    """原样读一个工作表的全部行, 返回 (工作表名, 行); sheet 为空取第一个. 工作表不存在或文件读不了抛 UnreadableFile."""
    names = sheet_names(file_name, data)
    name = sheet if sheet is not None else names[0]
    if name not in names:
        raise UnreadableFile(f"没有工作表 {name}")
    ext = extension(file_name)
    if ext == ".csv":
        return name, list(_csv_rows(data))
    if ext == ".xls":
        return name, list(_xls_rows(data, name))
    return name, list(_xlsx_rows(data, name))


def read_sheet(file_name: str, data: bytes, sheet: str | None = None) -> Sheet:
    """读一个工作表, 表头取第一个非空行; sheet 为空取第一个. 工作表不存在或文件读不了抛 UnreadableFile."""
    return sheet_from_rows(*raw_rows(file_name, data, sheet))


def sheet_from_rows(name: str, raw: Iterable[tuple[int, Sequence[Cell]]], header_row: int | None = None) -> Sheet:
    """原样的行 → 表头与数据行: 表头取行号为 header_row 的行, 为空时取第一个非空行; 表头之前的行不读, 之后跳过全空行;
    表头末尾的空单元格不算列. 指定的表头行不在工作表里或是空行时抛 MappingError (INVALID_VALUE)."""
    headers: list[str] | None = None
    found = 0
    rows: list[tuple[int, list[Cell]]] = []
    last = 0
    for number, cells in raw:
        last = number
        if headers is None:
            if header_row is not None and number != header_row:
                continue
            if all(is_blank(c) for c in cells):
                if header_row is not None:
                    raise MappingError("INVALID_VALUE", f"表头行 {header_row} 是空行")
                continue
            headers = [cell_text(c) for c in cells]
            while headers and not headers[-1]:
                headers.pop()
            found = number
            continue
        if all(is_blank(c) for c in cells):
            continue
        values = list(cells[: len(headers)])
        rows.append((number, values + [None] * (len(headers) - len(values))))
    if header_row is not None and headers is None:
        raise MappingError("INVALID_VALUE", f"表头行 {header_row} 超出工作表 (共 {last} 行)")
    headers = headers or []
    return Sheet(name=name, header_row=found, columns=unique_columns(headers), headers=headers, rows=rows)


def unique_columns(headers: Sequence[str]) -> list[str]:
    """各列的显示名: 空表头记 Column {位置}; 与前面重名的加 " (k)", k 从 2 起取第一个没被用过, 也不等于任何原始表头的,
    所以结果全表唯一 (如 name, name, name (2) → name, name (3), name (2))."""
    bases = [header or f"Column {index}" for index, header in enumerate(headers, start=1)]
    reserved = set(bases)
    used: set[str] = set()
    columns = []
    for base in bases:
        name, k = base, 2
        while name in used or (name != base and name in reserved):
            name, k = f"{base} ({k})", k + 1
        used.add(name)
        columns.append(name)
    return columns


def is_blank(cell: Cell) -> bool:
    """空单元格: 没有值或只有空白."""
    return cell is None or (isinstance(cell, str) and not cell.strip())


def _csv_rows(data: bytes) -> Iterator[tuple[int, list[Cell]]]:
    """CSV: 先按 UTF-8 (可带 BOM) 解码, 不行按 Windows-1252; 分隔符在逗号, 分号, 制表符, 竖线中嗅探."""
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("cp1252", errors="replace")
    try:
        dialect: type[csv.Dialect] | csv.Dialect = csv.Sniffer().sniff(text[:8192], delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text, newline=""), dialect)
    for row in reader:
        yield reader.line_num, list(row)


def _xlsx_book(data: bytes) -> Any:
    import openpyxl

    try:
        return openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # openpyxl 对坏文件抛多种异常
        raise UnreadableFile(f"读不了 XLSX 文件: {exc}") from exc


def _xlsx_rows(data: bytes, name: str) -> Iterator[tuple[int, list[Cell]]]:
    workbook = _xlsx_book(data)
    try:
        for number, row in enumerate(workbook[name].iter_rows(values_only=True), start=1):
            yield number, [c if isinstance(c, str | int | float | bool | date | datetime) else None for c in row]
    finally:
        workbook.close()


def _xls_book(data: bytes) -> Any:
    import xlrd

    try:
        return xlrd.open_workbook(file_contents=data)
    except Exception as exc:
        raise UnreadableFile(f"读不了 XLS 文件: {exc}") from exc


def _xls_rows(data: bytes, name: str) -> Iterator[tuple[int, list[Cell]]]:
    import xlrd

    book = _xls_book(data)
    sheet = book.sheet_by_name(name)
    for index in range(sheet.nrows):
        cells: list[Cell] = []
        for cell in sheet.row(index):
            if cell.ctype == xlrd.XL_CELL_DATE:
                cells.append(xlrd.xldate_as_datetime(cell.value, book.datemode))
            elif cell.ctype == xlrd.XL_CELL_BOOLEAN:
                cells.append(bool(cell.value))
            elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK, xlrd.XL_CELL_ERROR):
                cells.append(None)
            else:
                cells.append(cell.value)
        yield index + 1, cells


# ── 建议 ──


def suggest_mapping(kind: TableKind, columns: Sequence[str], *, shared: bool = True) -> dict[str, str | None]:
    """按列名为每个字段找一列: 先字段名本身 (含去掉 cents 的写法), 再按别名顺序, 最后是与别的表共有的列名
    (SHARED_ALIASES, shared 为假时不用); 一列只用一次."""
    normalized = {normalize(c): c for c in reversed(columns)}  # 规范化后同名的取靠前的列
    used: set[str] = set()
    mapping: dict[str, str | None] = {}
    common = SHARED_ALIASES.get(kind, {}) if shared else {}
    for field in SCHEMAS[kind].fields:
        own = normalize(field.name)
        candidates = [own, own.removesuffix("cents"), *ALIASES.get(field.name, ()), *common.get(field.name, ())]
        column = next((normalized[c] for c in candidates if c in normalized and normalized[c] not in used), None)
        if column is not None:
            used.add(column)
        mapping[field.name] = column
    return mapping


def suggest_kind(file_name: str, columns: Sequence[str]) -> TableKind | None:
    """按列名与文件名猜表种类: 必须字段 (与至少对应一个的字段组) 都能对应的种类中, 取能对应的字段个数最多
    (文件名提示加分) 的; 都对应不全为空. 与别的表共有的列名 (SHARED_ALIASES) 只在文件名有该表的提示时才算."""
    stem = normalize(PurePath(file_name).stem)
    best: tuple[float, TableKind] | None = None
    for kind, schema in SCHEMAS.items():
        hinted = any(h in stem for h in NAME_HINTS[kind])
        mapping = suggest_mapping(kind, columns, shared=hinted)
        if any(mapping[f.name] is None for f in schema.fields if f.required):
            continue
        group = ONE_OF.get(kind)
        if group and all(mapping[name] is None for name in group):
            continue
        matched = sum(1 for f in schema.fields if mapping[f.name] is not None)
        # 按能对应的字段个数 (不是占比): 完整的总账能对上的字段多于它的子集审批日志; 文件名提示加 2 分, 平局时取先列出的
        score = matched + (2 if hinted else 0)
        if best is None or score > best[0]:
            best = (score, kind)
    return best[1] if best else None


DATE_TEXT = re.compile(r"^\s*(\d{1,4})[/.-](\d{1,2})[/.-](\d{1,4})")


def suggest_date_format(values: Iterable[Cell]) -> DateFormat:
    """按样例文本日期猜写法: 年在前为 ymd; 第一段大于 12 为 dmy, 第二段大于 12 为 mdy; 定不了取 dmy (澳洲习惯)."""
    votes = {"dmy": 0, "mdy": 0, "ymd": 0}
    for value in values:
        if not isinstance(value, str):
            continue
        match = DATE_TEXT.match(value)
        if not match:
            continue
        first, second, _ = match.groups()
        if len(first) == 4:
            votes["ymd"] += 1
        elif int(first) > 12:
            votes["dmy"] += 1
        elif int(second) > 12:
            votes["mdy"] += 1
    best = max(votes, key=lambda k: votes[k])
    return cast(DateFormat, best) if votes[best] else "dmy"


def preview(import_id: str, file_name: str, sheets: list[str], sheet: Sheet) -> ImportPreview:
    """解析结果: 第一行数据作样例, 建议表种类, 列对应与日期写法."""
    first = sheet.rows[0][1] if sheet.rows else []
    sample = {column: cell_text(first[i]) if i < len(first) else "" for i, column in enumerate(sheet.columns)}
    kind = suggest_kind(file_name, sheet.columns)
    mapping = suggest_mapping(kind, sheet.columns) if kind is not None else {}
    date_columns = [
        sheet.columns.index(column)
        for field in (SCHEMAS[kind].fields if kind is not None else [])
        if field.type in ("date", "datetime") and (column := mapping.get(field.name)) is not None
    ] or list(range(len(sheet.columns)))
    values = [cells[i] for _, cells in sheet.rows[:200] for i in date_columns]
    return ImportPreview(
        import_id=import_id,
        file_name=file_name,
        sheets=sheets,
        sheet=sheet.name,
        row_count=len(sheet.rows),
        columns=sheet.columns,
        sample=sample,
        suggested_kind=kind,
        suggested_mapping=mapping,
        suggested_date_format=suggest_date_format(values),
    )


# ── 发给大模型的表格样本的遮蔽 (需求 R276) ──

# 人名一类的列 (取值要遮蔽), 表头先规整 (小写, 只留字母与数字, 即去掉 *, 空格与标点; 末尾的序号不计). 宁可多遮:
# Xero 日记账报表的 Posted By 是导出账号本人的姓名 (计划决定 12), 联系人导出的 *ContactName, FirstName, LastName,
# PersonNFirstName, POAttentionTo, SkypeName 等是联系人的姓名或账号.
# 两种认法: 恰为人名表头的单元格 (PERSONAL_HEADERS 之一, 或 "某某人" 列: 过去分词加 by, 如 postedby, createdby,
# approvedby, signedby, reviewedby, 另有 paidby, madeby) 不论在哪一行都认 (年份或日期作列名, 只有一列, 首格为空的
# 表头也认得); 只是含人名片段的 (PERSONAL_PARTS 之一, 或以 by 结尾) 要所在行像表头 (header_like), 免得数据行里的
# 科目名 Telephone & Internet 把整列遮掉
PERSONAL_HEADERS = frozenset(
    {
        "user",
        "username",
        "userid",
        "email",
        "emailaddress",
        "contactname",
        "contactperson",
        "firstname",
        "lastname",
        "fullname",
        "phone",
        "mobile",
        "attentionto",
        "skypename",
    }
)
BY_HEADER = re.compile(r"[a-z]*edby|paidby|madeby")  # "某某人" 列的完整写法
PERSONAL_PARTS = (
    "contactname",
    "firstname",
    "lastname",
    "fullname",
    "person",
    "user",
    "email",
    "phone",
    "mobile",
    "attentionto",
    "skype",
)
MASK = "[masked]"  # 遮蔽后的取值
EMAIL = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")  # 邮箱地址, 任何列里都遮蔽
# 数值或 ISO 日期的单元格文本 (金额, 代码, 序号; Excel 日期单元格经 cell_text 写成 ISO 日期): 有这种单元格的行不像表头.
# 文字写的日期 (如 30/06/2026) 不算, 只有文字的数据行因此会被当作表头, 方向是多遮
VALUE_TEXT = re.compile(r"^\(?[-+]?\$?[\d,]*\.?\d+\)?%?$|^\d{4}-\d{2}-\d{2}")


def header_key(header: str) -> str:
    """表头规整: 小写, 只留字母与数字, 去掉末尾的序号."""
    return normalize(header).rstrip("0123456789")


def personal_header(header: str) -> bool:
    """单元格是否恰为人名表头 (见 PERSONAL_HEADERS, BY_HEADER): 不论所在行的样子都认."""
    key = header_key(header)
    return key in PERSONAL_HEADERS or BY_HEADER.fullmatch(key) is not None


def personal_column(header: str) -> bool:
    """表头是否人名一类的列: 恰为人名表头, 或含人名片段 (PERSONAL_PARTS), 或以 by 结尾."""
    key = header_key(header)
    return bool(key) and (personal_header(header) or any(part in key for part in PERSONAL_PARTS) or key.endswith("by"))


def header_like(texts: Sequence[str]) -> bool:
    """一行是否像表头: 至少两个非空单元格, 且没有数值或 ISO 日期 (见 VALUE_TEXT). 只用来认含人名片段的表头."""
    filled = [t for t in texts if t]
    return len(filled) >= 2 and not any(VALUE_TEXT.match(t) for t in filled)


def mask_sample(rows: Iterable[Sequence[str]]) -> list[list[str]]:
    """表格样本的遮蔽, 发给大模型的表格文字 (DOC-SHEET 的样本行, DOC-CLASSIFY 的首页文字) 都先经它: rows 为依次各行的
    单元格文本 (含标题行; 表头用原始文本, 不用加了 " (2)" 的显示名). 一行里有人名表头 (恰为人名表头的不论行的样子,
    含人名片段的要行像表头, 见 PERSONAL_HEADERS 处的说明) 时, 之后各行该列的非空取值遮蔽为 MASK (表头本身照留, 大模型
    据此知道有这一列); 任何单元格里的邮箱地址遮蔽为 MASK. 行与单元格的个数不变."""
    personal: set[int] = set()
    masked = []
    for texts in rows:
        masked.append([MASK if i in personal and t else EMAIL.sub(MASK, t) for i, t in enumerate(texts)])
        like = header_like(texts)
        personal |= {i for i, t in enumerate(texts) if personal_header(t) or (like and personal_column(t))}
    return masked


# ── 单元格转换 ──


class CellError(ValueError):
    """一个单元格转换不了."""


def cell_text(cell: Cell) -> str:
    """单元格的文本: 去首尾空白; 整数值的浮点数去掉 .0 (Excel 把科目代码存成数值); 日期写 ISO."""
    if cell is None:
        return ""
    if isinstance(cell, bool):
        return "true" if cell else "false"
    if isinstance(cell, float):
        return str(int(cell)) if cell.is_integer() else repr(cell)
    if isinstance(cell, datetime):
        return cell.isoformat(sep=" ") if cell.time() != time(0) else cell.date().isoformat()
    if isinstance(cell, date):
        return cell.isoformat()
    return str(cell).strip()


def parse_amount(cell: Cell, cents: bool) -> int | None:
    """金额 → 整数分; 空为 None. cents 为真时值须是整数分, 否则按元读并四舍五入到分."""
    if isinstance(cell, bool):
        raise CellError("不是金额")
    negative = False
    if isinstance(cell, int | float):
        number = Decimal(repr(cell)) if isinstance(cell, float) else Decimal(cell)
    else:
        text = cell_text(cell)
        if not text:
            return None
        negative = (text.startswith("(") and text.endswith(")")) or text.endswith("-")
        text = text.strip("()").rstrip("-").replace(",", "").replace(" ", "")
        text = re.sub(r"^(AUD|A\$|\$)", "", text, flags=re.IGNORECASE).replace("$", "")
        try:
            number = Decimal(text)
        except InvalidOperation as exc:
            raise CellError("不是金额") from exc
    # 构造 Decimal 不受精度与指数范围限制, 之后的运算 (含取负) 才受; 所以先挡住非有限值与超大的数
    if not number.is_finite():
        raise CellError("不是金额")
    if number != 0 and number.adjusted() > 20:
        raise CellError("金额超出范围")
    try:
        if cents and number != number.to_integral_value():
            raise CellError("按分读的金额须为整数")
        value = int(number if cents else (number * 100).quantize(Decimal(1), rounding=ROUND_HALF_UP))
    except DecimalException as exc:
        raise CellError("金额超出范围") from exc
    value = -value if negative else value
    if abs(value) > MAX_CENTS:
        raise CellError("金额超出范围")
    return value


MONTH_NAMES = [
    "january",
    "february",
    "march",
    "april",
    "may",
    "june",
    "july",
    "august",
    "september",
    "october",
    "november",
    "december",
]
# 月份名: 全名, 三字母缩写, 另认 sept; 忽略大小写
MONTHS: dict[str, int] = {
    **{name: n for n, name in enumerate(MONTH_NAMES, start=1)},
    **{name[:3]: n for n, name in enumerate(MONTH_NAMES, start=1)},
    "sept": 9,
}
# 一个日期 (可带时间与时区) 的完整写法; 整个单元格必须匹配, 多出的内容算失败
DATE_TIME = re.compile(
    r"^\s*(?:"
    r"(?P<iy>\d{4})[/.-](?P<im>\d{1,2})[/.-](?P<id>\d{1,2})"  # 年在前: 总按年月日
    r"|(?P<a>\d{1,2})[/.-](?P<b>\d{1,2})[/.-](?P<c>\d{1,4})"  # 三段数字: 按 date_format, 年份位数在 _day 里查
    r"|(?P<nd>\d{1,2})[\s/.-]+(?P<nm>[A-Za-z]+)\.?[\s/.-]+(?P<ny>\d{2}|\d{4})"  # 日 月份名 年
    r")"
    r"(?:[T\s]+(?P<h>\d{1,2}):(?P<mi>\d{2})(?::(?P<s>\d{2})(?:\.(?P<f>\d+))?)?\s*(?P<ap>[AaPp][Mm])?)?"
    r"\s*(?P<tz>Z|[+-]\d{2}:?\d{2})?\s*$"
)


def _year(text: str) -> int:
    """两位年份按 20xx."""
    year = int(text)
    return year + 2000 if len(text) == 2 else year


@dataclass(frozen=True)
class Moment:
    """解析出的日期, 时间 (没有时为零点) 与时区 (没有为空)."""

    day: date  # 日期
    clock: time  # 时间
    zone: timezone | None  # 写明的时区


def _day(g: Mapping[str, str | None], date_format: DateFormat) -> date:
    """按匹配到的日期部分组出日期; 日期不存在, 月份名不认识, 写法与 date_format 不符抛 CellError."""
    try:
        if g["iy"] is not None:
            return date(int(g["iy"]), int(cast(str, g["im"])), int(cast(str, g["id"])))
        if g["a"] is not None:
            a, b, c = g["a"], cast(str, g["b"]), cast(str, g["c"])
            if date_format == "ymd":  # 年 (两位) 在前, 日在后; 四位年在前的已由上一种写法匹配
                if len(a) != 2 or len(c) > 2:
                    raise CellError("与写法 (年月日) 不符")
                return date(_year(a), int(b), int(c))
            if len(c) not in (2, 4):
                raise CellError("年份须为两位或四位")
            if date_format == "dmy":
                return date(_year(c), int(b), int(a))
            return date(_year(c), int(a), int(b))
        month = MONTHS.get(cast(str, g["nm"]).lower())
        if month is None:
            raise CellError("不是月份名")
        return date(_year(cast(str, g["ny"])), month, int(cast(str, g["nd"])))
    except ValueError as exc:
        raise CellError("日期不存在") from exc


def _moment(text: str, date_format: DateFormat) -> Moment:
    """解析一个日期 (可带时间与时区) 文本; 写法不符, 日期或时间不存在, 时区越界都抛 CellError."""
    match = DATE_TIME.match(text)
    if match is None:
        raise CellError("不是日期")
    g = match.groupdict()
    day = _day(g, date_format)
    clock = time(0)
    if g["h"] is not None:
        hour, minute, second = int(g["h"]), int(cast(str, g["mi"])), int(g["s"] or 0)
        if g["ap"] is not None:
            if not 1 <= hour <= 12:
                raise CellError("12 小时制的小时须为 1-12")
            hour = hour % 12 + (12 if g["ap"].lower() == "pm" else 0)
        micro = int((g["f"] or "")[:6].ljust(6, "0"))  # 小数秒保留到微秒, 更细的截去
        try:
            clock = time(hour, minute, second, micro)
        except ValueError as exc:
            raise CellError("时间不存在") from exc
    zone = None
    if g["tz"] == "Z":
        zone = UTC
    elif g["tz"] is not None:
        digits = g["tz"][1:].replace(":", "")
        hours, minutes = int(digits[:2]), int(digits[2:])
        if minutes > 59 or hours * 60 + minutes > 14 * 60:  # 时区偏移不超过 ±14:00
            raise CellError("时区不存在")
        offset = timedelta(hours=hours, minutes=minutes)
        zone = timezone(-offset if g["tz"].startswith("-") else offset)
    return Moment(day, clock, zone)


def parse_date(cell: Cell, date_format: DateFormat) -> date | None:
    """日期; 空为 None. 年在前的总按年月日, 三段数字按 date_format, 也认 "1 Jul 2025" 这类月份名; 年份两位或四位.
    后面可带时间与时区 (只取日期部分, 但时间与时区也须合法); 整个单元格须是这样的写法, 多出的内容算失败."""
    if isinstance(cell, datetime):
        return cell.date()
    if isinstance(cell, date):
        return cell
    text = cell_text(cell)
    return _moment(text, date_format).day if text else None


def parse_datetime(cell: Cell, date_format: DateFormat) -> datetime | None:
    """日期时间, 存 UTC; 空为 None. 写法同 parse_date, 时间可带秒 (小数秒保留到微秒) 与 AM/PM (此时小时须为 1-12);
    带时区 (Z, +10:00, +1000, 不超过 ±14:00) 的按其时区, 不带的按悉尼时间, 夏令时拨快时跳过的时刻
    (不存在的本地时间) 算失败; 只有日期时取零点."""
    if isinstance(cell, datetime):
        naive, zone = cell.replace(tzinfo=None), cell.tzinfo
    elif isinstance(cell, date):
        naive, zone = datetime.combine(cell, time(0)), None
    else:
        text = cell_text(cell)
        if not text:
            return None
        moment = _moment(text, date_format)
        naive, zone = datetime.combine(moment.day, moment.clock), moment.zone
    if zone is not None:
        return naive.replace(tzinfo=zone).astimezone(UTC)
    local = naive.replace(tzinfo=SYDNEY)
    if local.astimezone(UTC).astimezone(SYDNEY).replace(tzinfo=None) != naive:
        raise CellError("悉尼没有这个时刻 (夏令时拨快时跳过)")
    return local.astimezone(UTC)


TRUE = {"true", "t", "yes", "y", "1"}
FALSE = {"false", "f", "no", "n", "0"}


def parse_bool(cell: Cell) -> bool | None:
    """是否; 空为 None. 认 true/false, yes/no, y/n, 1/0 (忽略大小写)."""
    if isinstance(cell, bool):
        return cell
    text = cell_text(cell).lower()
    if not text:
        return None
    if text in TRUE:
        return True
    if text in FALSE:
        return False
    raise CellError("不是是否值")


def parse_integer(cell: Cell) -> int | None:
    """整数 (如发薪周期的天数); 空为 None. 认整数, 小数部分为 0 的数值单元格与可带千分位逗号的整数文本; 是否值,
    带小数的数与其他文本算失败."""
    if isinstance(cell, bool):
        raise CellError("不是整数")
    if isinstance(cell, int):
        return cell
    if isinstance(cell, float):
        if cell.is_integer():
            return int(cell)
        raise CellError("不是整数")
    text = cell_text(cell).replace(",", "")
    if not text:
        return None
    if not re.fullmatch(r"[+-]?[0-9]+", text):
        raise CellError("不是整数")
    return int(text)


# 科目类别与试算表期间的常见写法 → 系统取值; 其余报解析失败
ACCOUNT_TYPES = {
    "asset": "asset",
    "assets": "asset",
    "liability": "liability",
    "liabilities": "liability",
    "equity": "equity",
    "revenue": "revenue",
    "income": "revenue",
    "sales": "revenue",
    "expense": "expense",
    "expenses": "expense",
}
PERIODS = {
    "current": "current",
    "cy": "current",
    "currentyear": "current",
    "thisyear": "current",
    "prior": "prior",
    "py": "prior",
    "prioryear": "prior",
    "lastyear": "prior",
    "previous": "prior",
    "previousyear": "prior",
}


def parse_text(cell: Cell) -> str | None:
    text = cell_text(cell)
    return text or None


def _choice(values: Mapping[str, str], what: str) -> Callable[[Cell], str | None]:
    def parse(cell: Cell) -> str | None:
        text = normalize(cell_text(cell))
        if not text:
            return None
        if text not in values:
            raise CellError(f"不是{what}")
        return values[text]

    return parse


SPECIAL: dict[tuple[TableKind, str], Callable[[Cell], str | None]] = {
    (TableKind.accounts, "type"): _choice(ACCOUNT_TYPES, "科目类别 (资产, 负债, 权益, 收入, 费用)"),
    (TableKind.trial_balance, "period"): _choice(PERIODS, "期间 (current 或 prior)"),
}


# ── 整表转换 (按列对应或规整方案套用, 见 services/table_plan) ──


def convert_cell(kind: TableKind, field: TableField, cell: Cell, cents: bool, date_format: DateFormat) -> Any:
    """一个单元格按字段类型转换: 金额在 cents 为真时按整数分读, 否则按元; 科目类别与试算表期间按 SPECIAL 的写法.
    转换不了抛 CellError."""
    special = SPECIAL.get((kind, field.name))
    if special is not None:
        return special(cell)
    if field.type == "amount":
        return parse_amount(cell, cents=cents)
    if field.type == "date":
        return parse_date(cell, date_format)
    if field.type == "datetime":
        return parse_datetime(cell, date_format)
    if field.type == "bool":
        return parse_bool(cell)
    if field.type == "integer":
        return parse_integer(cell)
    return parse_text(cell)


def _line_id(values: Mapping[str, Any], counters: dict[str, int]) -> str:
    """分录行: {凭证号}-{该凭证内的行序号}."""
    entry = str(values.get("entry_no") or "")
    counters[entry] = counters.get(entry, 0) + 1
    return f"{entry}-{counters[entry]}"


def _bank_line_id(values: Mapping[str, Any], counters: dict[str, int]) -> str:
    """银行流水: {账户}-{该账户内的序号, 5 位}."""
    account = str(values.get("bank_account_id") or "")
    counters[account] = counters.get(account, 0) + 1
    return f"{account}-{counters[account]:05d}"


# 编号列没对应时的生成规则 (按文件中的顺序, 同一文件重导结果相同)
GENERATED_IDS: dict[TableKind, Callable[[Mapping[str, Any], dict[str, int]], str]] = {
    TableKind.journal_entries: _line_id,
    TableKind.sales_invoices: lambda values, _: str(values.get("invoice_no") or ""),
    TableKind.supplier_bills: lambda values, _: str(values.get("bill_no") or ""),
    TableKind.delivery_notes: lambda values, _: str(values.get("delivery_no") or ""),
    TableKind.bank_transactions: _bank_line_id,
}
