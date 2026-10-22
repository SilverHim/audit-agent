"""需求 R330 (Issue #50): 底稿记录. 程序执行终态与决定之后刷新 (依据, 结论草稿, 缺项), 审计师改结论, 补充说明与签署,
依据变了改为须重新复核, 不阻塞 FIN-01 与 RPT-01, 锁定后不能改, 刷新失败只记日志.

两种走法: 演示配置的离线走通 (tests/test_demo_walkthrough.walk, 编排里的刷新照常发生) 看各底稿的版本历史; 假存储里
直接放 Xero 演示 REV-08 第 1 次执行的执行记录, 被测总体与例外 (tests.xero_fakes) 与 PLN-01/03 的通过决定, 经接口
做审计师的三种操作."""

from collections.abc import AsyncIterator
from datetime import date
from typing import Any

import httpx
import pytest
import pytest_asyncio

from backend.contracts.models import (
    ExceptionDisposition,
    GateDecision,
    GateId,
    GateVerdict,
    Judgment,
    ObjectRef,
    Phase,
    ProcedureRun,
    RunPopulation,
    TrailKind,
    Workpaper,
    WorkpaperStatus,
)
from backend.core.config import get_settings
from backend.orchestration.decisions import GraphDecisionSubmit
from backend.orchestration.runner import GraphRunner
from backend.services import workpapers
from backend.services.procedure_library import ProcedureLibrary
from backend.services.workpapers import SATISFIED, refresh
from tests.fakes import (
    DECIDED_AT,
    ENGAGEMENT,
    NUMBAT_AUDITOR,
    FakeEngagementStore,
    FakeOutputStore,
    FakeTrail,
    acc01_decision,
    api_client,
    finished,
    make_procedure,
    make_run,
    make_runner,
    wait_idle,
)
from tests.fixture_sets import XERO
from tests.test_demo_walkthrough import Session, Walk, walk
from tests.xero_fakes import XERO17_MATERIALITY, xero_rev08_exception, xero_rev08_populations
from tests.xero_flow import demo_offline
from tests.xero_harness import LayeredMockLLMClient

EID = ENGAGEMENT.id
BASE = f"/api/v1/engagements/{EID}/workpapers"
REV08_PARAMS = {"window_days": 20, "period_end": date(2026, 9, 30)}  # 特别风险的窗口 (同 tests.xero_fakes)
# 验收 1: REV-08 第 1 次执行之后的结论草稿 (Xero 演示: 窗口内 6 张发票, 其中 5 张找得到收款, ORC1026 收款不符)
D3_DRAFT = {
    "zh": "范围: 期末前 20 天 (2026-09-10 至 2026-09-30) 入账的销售发票. 被测对象 6 项, 合计 AUD 9,900.00. "
    "入账期间正确: 6 项均符合. 有服务完成记录: 6 项均符合. 收款金额相符: 5 项中 1 项不符, 1 项不作判断. "
    "例外 1 条, 待处理 1 条.",
    "en": "Scope: Sales invoices posted in the 20 days before period end (2026-09-10 to 2026-09-30). "
    "6 items tested, totalling AUD 9,900.00. Recorded in the correct period: all 6 agree. "
    "Service completion record exists: all 6 agree. "
    "Receipts agree to the invoice: 1 of 5 do not agree, 1 not assessed. 1 exception, 1 pending.",
}
UNSIGNED = ["reviewer", "reviewed_on", "review_scope"]  # 未签署时缺的通用三项 (ASA 230 第 9(c) 段)


def test_every_required_item_has_a_check() -> None:
    """底稿配置的每个必填项都有判断是否已记录的规则, 规则表里也没有配置之外的代号."""
    library = ProcedureLibrary(get_settings().KNOWLEDGE_ROOT)
    assert {r.id for r in library.workpapers.required} == set(SATISFIED)


# ── 走通里的刷新 (验收 1, 2, 6, 7) ──

shared_loop = pytest.mark.asyncio(loop_scope="module")


@pytest_asyncio.fixture(scope="module", loop_scope="module")
async def walked() -> AsyncIterator[tuple[Walk, httpx.AsyncClient]]:
    """演示配置离线走通一次 (同 tests/test_demo_walkthrough 的 walked, 工作语言 en), 走完不关接口客户端."""
    XERO.ready()
    llm = LayeredMockLLMClient([XERO.runtime])
    async with demo_offline(llm) as (runner, service, client):
        yield await walk(Session(client, runner, service), llm), client


async def history(w: Walk, ref: str) -> list[Workpaper]:
    return await w.x.runner.deps.output_store.history(w.x.eid, Workpaper, f"WP-{ref}")


@pytest.mark.external
@shared_loop
async def test_rule_workpaper_after_first_run(walked: tuple[Walk, httpx.AsyncClient]) -> None:
    """验收 1: REV-08 第 1 次执行之后 WP-D-3 第 1 版为 draft, 依据为执行记录, 被测总体与例外, 结论草稿按模板; 之后
    ORC1026 的判断确认, 例外归因 (第 2 版), 草稿改写为已由判断解释."""
    w, _ = walked
    first, second, *rest = await history(w, "D-3")
    assert (first.version, first.status, first.procedure_id, first.area) == (1, WorkpaperStatus.draft, "REV-08", "D")
    assert first.basis == [
        ObjectRef(kind="AuditException", id="X-REV-08-1-R03-1", version=1),
        ObjectRef(kind="ProcedureRun", id="RUN-REV-08-1", version=1),
        ObjectRef(kind="RunPopulation", id="POP-RUN-REV-08-1", version=1),
    ]
    assert first.conclusion_draft is not None and first.conclusion_draft.model_dump() == D3_DRAFT
    assert first.missing == UNSIGNED
    assert ObjectRef(kind="AuditException", id="X-REV-08-1-R03-1", version=2) in second.basis
    assert second.conclusion_draft is not None
    assert second.conclusion_draft.zh.endswith("例外 1 条, 待处理 0 条, 已由判断解释 1 条.")
    assert rest == []


@pytest.mark.external
@shared_loop
async def test_judgment_workpaper_follows_reviews(walked: tuple[Walk, httpx.AsyncClient]) -> None:
    """验收 2: REV-03 的 WP-D-2 第 1 版列 3 份合同的建议与置信度, 待复核 3 项, 末句同 Q27; 三个复核决定各出一版,
    最后一版待复核 0 项."""
    w, _ = walked
    versions = await history(w, "D-2")
    drafts = [v.conclusion_draft.zh if v.conclusion_draft else "" for v in versions]
    assert drafts[0] == (
        "判断对象 3 个: CON-X01 建议接受, 置信度 0.84, 待复核; CON-X02 建议接受, 置信度 0.84, 待复核; "
        "CON-X03 建议接受, 置信度 0.82, 待复核. 待复核 3 项. 最终结论待审计师复核后确定."
    )
    assert [d.split("待复核 ")[-1][0] for d in drafts] == ["3", "2", "1", "0"]
    assert all(d.endswith("最终结论待审计师复核后确定.") for d in drafts)


@pytest.mark.external
@shared_loop
async def test_unsigned_workpapers_do_not_block_and_lock_after_archive(walked: tuple[Walk, httpx.AsyncClient]) -> None:
    """验收 6, 7: 走通时没有签署任何底稿, FIN-01 与 RPT-01 照常通过到归档; 列表按档案顺序, 各程序底稿都有记录且都是
    draft; 归档 (锁定) 后三种操作都报 423 ENGAGEMENT_LOCKED, 刷新不再写."""
    w, client = walked
    base = f"/api/v1/engagements/{w.x.eid}/workpapers"
    listed = (await client.get(base)).json()
    assert [r["ref"] for r in listed] == [
        "00",
        *("A-1", "A-2", "A-3", "A-4", "B-1", "B-2", "C", "C-1", "C-2", "C-3", "C-4", "D", "D-1", "D-2", "D-3", "D-4"),
        *("G-1", "G-2", "X", "X-A", "X-B", "X-C", "Z-1", "Z-2", "Z-3"),
        "S",
    ]
    assert {r["status"] for r in listed} == {"draft"}
    posts = [
        ("D-3/conclusion", {"text": "x"}),
        ("D-3/notes", {"text": "x"}),
        ("D-3/signoff", {"scope": "full"}),
    ]
    for path, body in posts:
        response = await client.post(f"{base}/{path}", json=body)
        assert (response.status_code, response.json()["detail"]["code"]) == (423, "ENGAGEMENT_LOCKED")
    assert await refresh(w.x.runner.deps, w.x.eid) == []


# ── 审计师的三种操作 (验收 3, 4, 5, 7) ──


def rev08_run(run_no: int) -> ProcedureRun:
    return finished(make_run("REV-08", run_no, dict(REV08_PARAMS)))


def rev08_outputs(run: ProcedureRun) -> list[Any]:
    """REV-08 一次执行的执行记录, 被测总体与例外 (Xero 演示)."""
    population = RunPopulation(id=f"POP-{run.id}", run_id=run.id, procedure_id="REV-08", rules=xero_rev08_populations())
    return [run, population, xero_rev08_exception(run)]


PLN_APPROVE = GateDecision(
    id="D-PLN",
    gate=GateId.PLN_01_03,
    target_id="MATERIALITY",
    target_version=1,
    decision=GateVerdict.approve,
    materiality=XERO17_MATERIALITY,
    decided_by=NUMBAT_AUDITOR.id,
    decided_at=DECIDED_AT,
)


async def seeded() -> tuple[GraphRunner, FakeOutputStore]:
    """程序包有 REV-08 与 PLN-01, Store 里有 REV-08 第 1 次执行的产出与 PLN-01/03 的通过决定 (没写理由), 刷新一次:
    WP-D-3 与 WP-A-2 各第 1 版."""
    runner, store = make_runner()
    pack = [make_procedure("REV-08", Phase.P2_Executing), make_procedure("PLN-01", Phase.P1_Planning)]
    await store.put(EID, [*pack, *rev08_outputs(rev08_run(1)), PLN_APPROVE], actor="A6")
    assert {w.id for w in await refresh(runner.deps, EID)} == {"WP-A-2", "WP-D-3", "WP-S", "WP-00"}
    return runner, store


def decisions(runner: GraphRunner) -> list[dict[str, Any]]:
    trail = runner.deps.trail
    assert isinstance(trail, FakeTrail)
    return [e.payload | {"procedure_id": e.procedure_id} for e in trail.entries if e.kind is TrailKind.decision]


async def test_conclusion_and_notes() -> None:
    """验收 3: 改结论出新版本 (草稿照留), decision 留痕一行; 补充说明补齐 A-2 的选择理由 (materiality_factors, 可由
    补充说明补齐) 后它从 missing 去掉; 不可补齐的项, 不认识的项与空文字被拒, 不写."""
    runner, store = await seeded()
    async with api_client(runner) as client:
        before = await store.get(EID, Workpaper, "WP-A-2")
        assert before is not None and "materiality_factors" in before.missing
        concluded = await client.post(f"{BASE}/D-3/conclusion", json={"text": " Agreed, see D-4 for ORC1026. "})
        noted = await client.post(
            f"{BASE}/A-2/notes", json={"text": "Revenue is the stable benchmark.", "item": "materiality_factors"}
        )
        rejected = [
            await client.post(f"{BASE}/A-2/notes", json={"text": "x", "item": "reviewer"}),  # 签署才能补
            await client.post(f"{BASE}/A-2/notes", json={"text": "x", "item": "nope"}),
            await client.post(f"{BASE}/A-2/notes", json={"text": "  "}),
            await client.post(f"{BASE}/D-3/conclusion", json={"text": ""}),
        ]
    assert concluded.status_code == 200, concluded.text
    body = concluded.json()
    assert (body["version"], body["conclusion"], body["conclusion_by"]) == (2, "Agreed, see D-4 for ORC1026.", "U-1")
    assert body["conclusion_draft"]["en"] == D3_DRAFT["en"]
    assert noted.status_code == 200, noted.text
    assert (
        noted.json()["notes"][0]["item"] == "materiality_factors"
        and "materiality_factors" not in noted.json()["missing"]
    )
    codes = [(r.status_code, r.json()["detail"]["code"]) for r in rejected]
    assert codes == [(422, "INVALID_VALUE"), (422, "INVALID_VALUE"), (422, "MISSING_FIELD"), (422, "MISSING_FIELD")]
    assert [(d["action"], d["procedure_id"], d["version"]) for d in decisions(runner)] == [
        ("workpaper_conclusion", "REV-08", 2),
        ("workpaper_note", "PLN-01", 2),
    ]


async def test_signoff() -> None:
    """验收 4: 签署出新版本 status signed, signoff 记签署人, 时间, 范围与意见, 未签署缺的三项随之去掉, decision 留痕;
    意见选填; 范围不认识 422 (请求体校验); 同一人同样再签不出新版本."""
    runner, _ = await seeded()
    async with api_client(runner) as client:
        signed = await client.post(
            f"{BASE}/D-3/signoff", json={"scope": "full", "opinion": "Agreed to supporting records."}
        )
        again = await client.post(
            f"{BASE}/D-3/signoff", json={"scope": "full", "opinion": "Agreed to supporting records."}
        )
        plain = await client.post(f"{BASE}/A-2/signoff", json={"scope": "conclusion_only"})
        bad = await client.post(f"{BASE}/D-3/signoff", json={"scope": "glance"})
    assert signed.status_code == 200, signed.text
    body = signed.json()
    assert (body["version"], body["status"], body["missing"]) == (2, "signed", [])
    signoff = body["signoff"]
    assert (signoff["signed_by"], signoff["scope"], signoff["opinion"]) == (
        "U-1",
        "full",
        "Agreed to supporting records.",
    )
    assert again.json()["version"] == 2
    assert (plain.status_code, plain.json()["signoff"]["opinion"]) == (200, "")
    assert bad.status_code == 422
    assert [(d["action"], d["scope"], d["version"]) for d in decisions(runner)] == [
        ("workpaper_signoff", "full", 2),
        ("workpaper_signoff", "conclusion_only", 2),
    ]


def rerun(run: ProcedureRun) -> list[Any]:
    return rev08_outputs(rev08_run(2))


def attributed(run: ProcedureRun) -> list[Any]:
    exception = xero_rev08_exception(run)
    update = {"version": 2, "disposition": ExceptionDisposition.attributed, "attributed_by": "J-REV-02-1-ORC1026"}
    return [exception.model_copy(update=update)]


@pytest.mark.parametrize("change", [rerun, attributed], ids=["rerun", "exception_decided"])
async def test_signed_workpaper_needs_rereview_when_basis_changes(change: Any) -> None:
    """验收 5: 已签署的 WP-D-3 在 REV-08 重跑 (RUN-REV-08-2) 或其例外有新决定 (归因, 第 2 版) 之后刷新出新版本,
    status needs_rereview, last_signoff 为原签署, signoff 清空; 依据没变的刷新不出新版本."""
    runner, store = await seeded()
    async with api_client(runner) as client:
        signed = (await client.post(f"{BASE}/D-3/signoff", json={"scope": "sample"})).json()
    assert await refresh(runner.deps, EID) == []  # 依据没变
    await store.put(EID, change(rev08_run(1)), actor="A6")
    [refreshed] = [p for p in await refresh(runner.deps, EID) if p.ref == "D-3"]
    assert (refreshed.id, refreshed.version, refreshed.status) == ("WP-D-3", 3, WorkpaperStatus.needs_rereview)
    assert refreshed.signoff is None and refreshed.last_signoff is not None
    assert refreshed.last_signoff.model_dump(mode="json") == signed["signoff"]
    assert refreshed.missing == UNSIGNED
    assert await refresh(runner.deps, EID) == []


async def test_edit_after_signoff_needs_rereview() -> None:
    """签署后改结论也算内容变了: 新版本 needs_rereview, 签署移到 last_signoff."""
    runner, _ = await seeded()
    async with api_client(runner) as client:
        await client.post(f"{BASE}/D-3/signoff", json={"scope": "full"})
        edited = (await client.post(f"{BASE}/D-3/conclusion", json={"text": "Revised."})).json()
    assert (edited["status"], edited["signoff"], edited["last_signoff"]["scope"]) == ("needs_rereview", None, "full")


@pytest.mark.parametrize(
    ("setup", "path", "code"),
    [
        ({"locked": True}, "D-3/conclusion", (423, "ENGAGEMENT_LOCKED")),
        ({"locked": True}, "D-3/notes", (423, "ENGAGEMENT_LOCKED")),
        ({"locked": True}, "D-3/signoff", (423, "ENGAGEMENT_LOCKED")),
        ({}, "C-3/conclusion", (404, "NOT_FOUND")),  # CSH-05 没有执行, 没有记录
        ({"busy": True}, "D-3/conclusion", (409, "ENGAGEMENT_BUSY")),
    ],
)
async def test_operations_rejected(setup: dict[str, bool], path: str, code: tuple[int, str]) -> None:
    """验收 7: 项目锁定 (签字或拒绝承接) 后三种操作 423; 没有记录的底稿 404; 项目锁被占 409. 都不写."""
    runner, store = await seeded()
    engagements = runner.deps.engagement_store
    assert isinstance(engagements, FakeEngagementStore)
    if setup.get("locked"):
        await engagements.update(EID, locked=True)
    if setup.get("busy"):
        runner.acquire(EID)
    body = {"scope": "full"} if path.endswith("signoff") else {"text": "x"}
    async with api_client(runner) as client:
        response = await client.post(f"{BASE}/{path}", json=body)
    assert (response.status_code, response.json()["detail"]["code"]) == code
    assert {w.version for w in await store.latest(EID, Workpaper)} == {1}


# ── 刷新失败只记日志 (验收 8) ──


async def test_refresh_failure_only_logs(monkeypatch: pytest.MonkeyPatch, caplog: pytest.LogCaptureFixture) -> None:
    """刷新出错 (构造的异常) 时程序执行与图照常推进, 只记日志; ACC-01 决定之后的刷新补上 B-2 (ACC-04) 与 B-1."""
    runner, store = make_runner()
    real = workpapers.load

    async def broken(*args: Any, **kwargs: Any) -> Any:
        raise RuntimeError("构造的刷新失败")

    monkeypatch.setattr(workpapers, "load", broken)
    await runner.start(EID)
    await wait_idle(runner, EID)
    assert runner.run_error(EID) is None and await store.latest(EID, Workpaper) == []
    assert "刷新底稿记录失败" in caplog.text
    monkeypatch.setattr(workpapers, "load", real)
    await GraphDecisionSubmit(runner.app, runner.deps, runner).submit(EID, acc01_decision(GateVerdict.approve))
    await wait_idle(runner, EID)
    assert {"WP-B-1", "WP-B-2"} <= {w.id for w in await store.latest(EID, Workpaper)}


@pytest.mark.external
@shared_loop
async def test_latest_judgment_uses_numeric_run_order(walked: tuple[Walk, httpx.AsyncClient]) -> None:
    """第 10 次执行的判断取代第 9 次, 不按执行 ID 的字符串大小选旧判断, 也不混入人工判断."""
    w, _ = walked
    [original, *_] = await w.x.runner.deps.output_store.latest(w.x.eid, Judgment)
    ninth, tenth = finished(make_run(original.procedure_id, 9)), finished(make_run(original.procedure_id, 10))
    old = original.model_copy(update={"id": "J-9", "run_id": ninth.id, "version": 3})
    new = original.model_copy(update={"id": "J-10", "run_id": tenth.id})
    human = new.model_copy(update={"id": "J-10-H", "version": 2})
    assert workpapers.model_judgments([new, human, old], [tenth, ninth]) == [new]


async def test_currency_change_requires_rereview_and_preserves_prior_draft() -> None:
    """币种改后新底稿显示新代码并要求重新复核, 旧版与原始金额不改."""
    runner, store = await seeded()
    original = await store.get(EID, Workpaper, "WP-D-3")
    async with api_client(runner) as client:
        await client.post(f"{BASE}/D-3/signoff", json={"scope": "full"})
        response = await client.put(f"/api/v1/engagements/{EID}/currency", json={"currency": "NZD"})
        assert response.status_code == 200
    [changed] = [paper for paper in await refresh(runner.deps, EID) if paper.ref == "D-3"]
    assert changed.status is WorkpaperStatus.needs_rereview and changed.last_signoff is not None
    assert changed.conclusion_draft is not None and "NZD 9,900.00" in changed.conclusion_draft.zh
    assert await store.get(EID, Workpaper, "WP-D-3", 1) == original
    assert await refresh(runner.deps, EID) == []


@pytest.mark.external
@shared_loop
async def test_generic_workpaper_snapshots(walked: tuple[Walk, httpx.AsyncClient]) -> None:
    """共用走通核对普通规则与登记程序, 识别字段和检查结果进入快照, 不复制原始录入人员."""
    w, _ = walked
    papers = {ref: (await history(w, ref))[0] for ref in ("C-3", "A-4", "G-1", "B-2")}
    tables = {
        ref: next(s for s in paper.content.sections if s.id == "details").tables[0] for ref, paper in papers.items()
    }
    assert len(tables["C-3"].rows) == 2
    assert len(tables["B-2"].rows) == 3
    assert {row["status"] for row in tables["B-2"].rows} == {"✓"}
    assert len(tables["G-1"].rows) == 608
    assert {row["source_id"] for row in tables["A-4"].rows} >= {"090", "091", "840"}
    for paper in papers.values():
        snapshot = paper.content.model_dump_json()
        assert "{{" not in snapshot
    for data in tables.values():
        assert all(not ({"created_by", "approved_by", "posted_by"} & row.keys()) for row in data.rows)


@pytest.mark.external
@shared_loop
async def test_judgment_and_extraction_snapshots(walked: tuple[Walk, httpx.AsyncClient]) -> None:
    """同一走通核对三份合同十一条原文、逐对象推理与复核, 两份银行回函的抽取字段."""
    from io import BytesIO

    from openpyxl import load_workbook

    from backend.services.workpaper_xlsx import render

    w, _ = walked
    first, *_, last = await history(w, "D-2")
    details = next(s for s in last.content.sections if s.id == "details")
    assert [s.title.en for s in details.children] == ["CON-X01", "CON-X02", "CON-X03"]
    terms = [table for child in details.children for table in child.tables if table.id.startswith("terms_")]
    assert sum(len(table.rows) for table in terms) == 11
    assert all(row["quote"] and row["page"] >= 1 for table in terms for row in table.rows)
    assert len(details.tables[0].rows) == 3
    for child in details.children:
        assert child.generated_language == "en"
        reason = next(t for t in child.tables if t.id.startswith("reasoning_"))
        assert reason.rows and any(row["kind"].en == "Fact" and row["evidence"] for row in reason.rows)
        assert next(t for t in child.tables if t.id.startswith("confidence_")).rows[0]["confidence"] in {0.84, 0.82}
        assert next(t for t in child.tables if t.id.startswith("critic_")).rows
        reviews = next(t for t in child.tables if t.id.startswith("reviews_"))
        assert reviews.rows[0]["decision"].en == "Confirm" and reviews.rows[0]["reviewer"] == w.x.owner
    first_details = next(s for s in first.content.sections if s.id == "details")
    assert {r["status"].en for r in first_details.tables[0].rows} == {"Awaiting auditor decision"}
    extracted = (await history(w, "D-1"))[-1]
    contracts = next(s for s in extracted.content.sections if s.id == "details").children
    assert len(contracts) == 3 and sum(len(c.tables[1].rows) for c in contracts) == 11
    confirmation = (await history(w, "C-2"))[-1]
    fields = next(s for s in confirmation.content.sections if s.id == "details").tables[0]
    assert len(fields.rows) == 2
    assert {r["accrued_interest_cents"] for r in fields.rows} == {0, 4853}
    for language in ("zh", "en"):
        cells = [cell.value for row in load_workbook(BytesIO(render(last, language))).active for cell in row]
        assert ("判断文字为英文生成" in cells) == (language == "zh")
        assert all(row["quote"] in cells for table in terms for row in table.rows)


@pytest.mark.external
@shared_loop
async def test_judgment_missing_reasoning_and_extraction_changes(walked: tuple[Walk, httpx.AsyncClient]) -> None:
    """缺推理与指针明确进入缺项; 人工抽取修改保留字段、版本、人、时间与理由."""
    from dataclasses import replace

    from backend.contracts.data import ExtractionRecord
    from backend.services.workpaper_content import build_content
    from backend.services.workpapers import load, refreshed, view_of

    w, _ = walked
    snapshot = await load(w.x.runner.deps, w.x.eid)
    sources = snapshot.sources
    original = next(j for j in sources.judgments if j.procedure_id == "REV-03")
    changed = original.model_copy(update={"version": original.version + 1, "reasoning_chain": [], "evidence": []})
    broken = replace(sources, judgments=[changed if j.id == original.id else j for j in sources.judgments])
    paper = (await history(w, "D-2"))[-1]
    new = refreshed(paper, snapshot.templates["REV-03"], snapshot.required["REV-03"], broken)
    assert {"conclusions", "results"} <= set(new.missing)
    note = ExtractionRecord(
        target_type="contract",
        target_id="CON-X01",
        version=2,
        data={"amount_cents": 1500},
        field="amount_cents",
        reason="Agreed to signed amendment.",
        changed_by=w.x.owner,
        changed_at=DECIDED_AT,
    )
    edited = replace(sources, extractions=[*sources.extractions, note])
    contract_paper = (await history(w, "D-1"))[-1]
    content = build_content(view_of(snapshot.templates["DOC-04"], edited), edited, contract_paper)
    details = next(s for s in content.sections if s.id == "details")
    edits = next(t for c in details.children for t in c.tables if t.id == "extraction_changes" and t.rows)
    assert edits.rows == [
        {
            "source_id": "CON-X01",
            "version": 2,
            "field": "amount_cents",
            "before": None,
            "after": "AUD 15.00",
            "changed_by": w.x.owner,
            "changed_at": DECIDED_AT.isoformat(),
            "reason": "Agreed to signed amendment.",
        }
    ]
    assert contract_paper.content != content


@pytest.mark.external
@shared_loop
async def test_decision_workpapers_record_materials_and_calculations(walked: tuple[Walk, httpx.AsyncClient]) -> None:
    """复用走通核对重要性计算、风险、承接评分卡、计划修订与报告文本."""
    w, _ = walked
    papers = {ref: (await history(w, ref))[-1] for ref in ("A-1", "A-2", "A-3", "B-1", "Z-1", "Z-2", "Z-3")}
    details = {ref: next(s for s in p.content.sections if s.id == "details") for ref, p in papers.items()}
    [calculation] = next(t for t in details["A-2"].tables if t.id == "materiality_calculation").rows
    assert {
        key: calculation[key]
        for key in (
            "base_amount_cents",
            "pct",
            "overall_cents",
            "performance_pct",
            "performance_cents",
            "trivial_pct",
            "trivial_cents",
        )
    } == {
        "base_amount_cents": 3243100,
        "pct": 0.01,
        "overall_cents": 32431,
        "performance_pct": 0.5,
        "performance_cents": 16216,
        "trivial_pct": 0.05,
        "trivial_cents": 1622,
    }
    assert calculation["specific"].en == "Not applicable"
    risks = next(t for t in details["A-3"].tables if t.id == "risk_register").rows
    assert {r["id"]: r["risk"].en for r in risks} == {
        "RISK-REV-OCC": "Significant risk",
        "RISK-REV-CUT": "Significant risk",
        "RISK-CASH-EXI": "High",
    }
    assert len(next(t for t in details["B-1"].tables if t.id == "acceptance").rows) == 3
    assert next(t for t in details["B-1"].tables if t.id == "decision_history").rows[0]["decision"].en == "Approved"
    assert details["A-1"].children and all(child.paragraphs for child in details["A-1"].children)
    assert all(any(t.id == "plan_material" for t in child.tables) for child in details["A-1"].children)
    assert details["Z-2"].children and all(child.paragraphs for child in details["Z-2"].children)
    assert next(t for t in details["Z-1"].tables if t.id == "completion_material").rows
    assert next(t for t in details["Z-3"].tables if t.id == "decision_history").rows  # 退回也留作决定历史.


@pytest.mark.external
@shared_loop
async def test_materiality_reason_supplement_and_bound_candidate_version(
    walked: tuple[Walk, httpx.AsyncClient],
) -> None:
    """空理由明确提示且能补齐; 新候选不倒算旧决定, 两次选择分别使用各自绑定的版本."""
    from dataclasses import replace

    from backend.contracts.models import MaterialityCandidates, WorkpaperNote
    from backend.services.workpaper_content import build_content
    from backend.services.workpapers import load, missing_of, view_of

    w, _ = walked
    snapshot = await load(w.x.runner.deps, w.x.eid)
    paper = snapshot.workpapers["WP-A-2"]
    sources = snapshot.sources
    template = snapshot.templates["PLN-01"]
    decisions = [d.model_copy(update={"reason": None}) if d.gate is GateId.PLN_01_03 else d for d in sources.decisions]
    empty = replace(sources, decisions=decisions)
    view = view_of(template, empty)
    blank = build_content(view, empty, paper)
    row = next(t for s in blank.sections for t in s.tables if t.id == "materiality_calculation").rows[0]
    assert row["reason"].en == "Missing: no reason recorded at the decision"
    assert "materiality_factors" in missing_of(snapshot.required[template.id], view, empty, paper)
    noted = paper.model_copy(
        update={
            "notes": [
                WorkpaperNote(
                    text="Revenue is the stable benchmark.",
                    item="materiality_factors",
                    added_by=w.x.owner,
                    added_at=DECIDED_AT,
                )
            ]
        }
    )
    content = build_content(view, empty, noted)
    assert (
        next(t for s in content.sections for t in s.tables if t.id == "materiality_calculation").rows[0]["reason"]
        == "Revenue is the stable benchmark."
    )
    assert "materiality_factors" not in missing_of(snapshot.required[template.id], view, empty, noted)
    candidate = next(o for o in sources.outputs["materiality_candidates"] if isinstance(o, MaterialityCandidates))
    altered = candidate.model_copy(
        update={
            "version": 2,
            "candidates": [
                c.model_copy(update={"base_amount_cents": 5000000}) if c.basis == "revenue" else c
                for c in candidate.candidates
            ],
        }
    )
    changed_sources = replace(empty, outputs={**sources.outputs, "materiality_candidates": [altered]})
    content = build_content(view_of(template, changed_sources), changed_sources, paper)
    assert (
        next(t for s in content.sections for t in s.tables if t.id == "materiality_calculation").rows[0][
            "overall_cents"
        ]
        == 32431
    )

    from datetime import timedelta

    selected = next(d for d in decisions if d.gate is GateId.PLN_01_03)
    later = selected.model_copy(
        update={"id": "D-PLN-REVISED", "target_version": 2, "decided_at": selected.decided_at + timedelta(seconds=1)}
    )
    revised_sources = replace(
        changed_sources,
        decisions=[*decisions, later],
        materials={**sources.materials, (altered.id, altered.version): altered},
    )
    updated = workpapers.refreshed(paper, template, snapshot.required[template.id], revised_sources)
    table = next(t for s in updated.content.sections for t in s.tables if t.id == "materiality_calculation")
    assert [r["overall_cents"] for r in table.rows] == [32431, 50000]
    assert {(r.kind, r.id, r.version) for r in updated.basis} >= {
        ("MaterialityCandidates", "MATERIALITY", 1),
        ("MaterialityCandidates", "MATERIALITY", 2),
    }


@pytest.mark.external
@shared_loop
async def test_index_standards_and_misstatement_snapshots(walked: tuple[Walk, httpx.AsyncClient]) -> None:
    """总索引覆盖配置, 相对链接与导出文件名一致; 准则对照保留实际缺项; 错报表不把已更正额累计."""
    from io import BytesIO

    from openpyxl import load_workbook

    from backend.services.workpaper_xlsx import filename, render

    w, _ = walked
    papers = {p.ref: p for p in await w.x.runner.deps.output_store.latest(w.x.eid, Workpaper)}
    [index] = [t for s in papers["00"].content.sections for t in s.tables if t.id == "workpaper_index"]
    assert [row["ref"] for row in index.rows] == w.x.runner.deps.library.workpaper_refs()
    assert next(row for row in index.rows if row["ref"] == "B-3")["status"].en == "Feature disabled"
    assert next(row for row in index.rows if row["ref"] == "G-3")["status"].en == "Feature disabled"
    for language in ("zh", "en"):
        sheet = load_workbook(BytesIO(render(papers["00"], language))).active
        for ref, link in index.links.items():
            assert getattr(link, language) == filename(papers[ref], language)
            cell = next(cell for row in sheet for cell in row if cell.value == ref and cell.hyperlink)
            assert cell.hyperlink.target == filename(papers[ref], language)
    standards = next(t for s in papers["S"].content.sections for t in s.tables if t.id == "standards")
    assert len(standards.rows) == 15
    assert next(r for r in standards.rows if r["standard"] == "ASA 240")["paragraph"] == "45-48"
    asa330 = next(r for r in standards.rows if r["standard"] == "ASA 330")
    assert "lead schedules C and D" in asa330["remarks"].en
    assert next(r for r in standards.rows if r["standard"] == "ASA 600")["status"].en.startswith("Not applicable")
    details = next(s for s in papers["X"].content.sections if s.id == "details")
    errors = next(t for t in details.tables if t.id == "misstatements")
    assert sorted(r["amount_cents"] for r in errors.rows) == [4853, 75000]
    assert {r["status"].en for r in errors.rows} == {"Corrected"}
    assert {r["source_ref"] for r in errors.rows} == {"G-2", "C-4"}
    assert all(r["decided_by"] == w.x.owner and r["decided_at"] for r in errors.rows)
    assert sum(r["accumulated_cents"] for r in errors.rows) == 0
    sheet = load_workbook(BytesIO(render(papers["X"], "en"))).active
    formulas = [cell.value for row in sheet for cell in row if cell.data_type == "f"]
    assert len(formulas) == 3 and any(f.startswith("=SUM(") for f in formulas)
    assert sum(f.startswith("=IF(") for f in formulas) == 2


async def test_system_workpaper_operations_are_versioned_and_stable() -> None:
    """系统底稿使用同一签署与补充入口, 相关索引刷新之后再刷新不写, 旧版可以独立导出."""
    runner, store = await seeded()
    before = await store.get(EID, Workpaper, "WP-S")
    async with api_client(runner) as client:
        signed = await client.post(f"{BASE}/S/signoff", json={"scope": "full"})
        assert signed.status_code == 200 and signed.json()["status"] == "signed"
        noted = await client.post(
            f"{BASE}/00/notes", json={"text": "Scope limitations remain visible on each workpaper."}
        )
        assert noted.status_code == 200
    assert await refresh(runner.deps, EID) == []
    assert await store.get(EID, Workpaper, "WP-S", 1) == before
    assert (await store.get(EID, Workpaper, "WP-S")).status is WorkpaperStatus.signed


@pytest.mark.external
@shared_loop
async def test_summary_excludes_withdrawn_and_standards_show_missing_reason(
    walked: tuple[Walk, httpx.AsyncClient],
) -> None:
    """累计表陈旧时也不把撤回金额带回预览; ASA 320 缺理由与没有任何对应底稿的要求明确标出."""
    from dataclasses import replace

    from backend.contracts.models import MisstatementStatus, UncorrectedItem, UncorrectedSchedule
    from backend.services.workpaper_summaries import misstatement_section, standards_table
    from backend.services.workpapers import load, view_of

    w, _ = walked
    snapshot = await load(w.x.runner.deps, w.x.eid)
    sources = snapshot.sources
    uncorrected = [
        m.model_copy(
            update={
                "status": MisstatementStatus.uncorrected,
                "corrected_cents": 0,
                "residual_cents": None,
                "verified_by_judgment_id": None,
            }
        )
        for m in sources.misstatements
    ]
    original = next(o for o in sources.outputs["uncorrected_schedule"] if isinstance(o, UncorrectedSchedule))
    schedule = original.model_copy(
        update={
            "items": [
                UncorrectedItem(misstatement_id=m.id, amount_cents=m.amount_cents, direction=m.direction)
                for m in uncorrected
            ],
            "total_cents": 79853,
        }
    )
    altered = replace(
        sources, misstatements=uncorrected, outputs={**sources.outputs, "uncorrected_schedule": [schedule]}
    )
    details = misstatement_section(view_of(snapshot.templates["FIN-04"], altered), altered)
    assert next(t for t in details.tables if t.id == "misstatement_totals").rows[0]["total_cents"] == 79853
    withdrawn = uncorrected[0].model_copy(update={"status": MisstatementStatus.withdrawn})
    altered = replace(altered, misstatements=[withdrawn, *uncorrected[1:]])
    details = misstatement_section(view_of(snapshot.templates["FIN-04"], altered), altered)
    assert (
        next(t for t in details.tables if t.id == "misstatement_totals").rows[0]["total_cents"]
        == uncorrected[1].amount_cents
    )
    papers = {p.ref: p for p in snapshot.workpapers.values()}
    papers["A-2"] = papers["A-2"].model_copy(update={"missing": ["materiality_factors"]})
    table = standards_table(sources, papers)
    assert next(r for r in table.rows if r["standard"] == "ASA 320")["status"].en == "Missing selection rationale"
    assert next(r for r in table.rows if r["standard"] == "ASA 250")["status"].en == "Missing"


@pytest.mark.external
@shared_loop
async def test_cash_and_revenue_lead_schedules(walked: tuple[Walk, httpx.AsyncClient]) -> None:
    """同一 Xero 走通的导引表: 审定数只加已验证分录, 不重述上年数, 期间不可比显式提示."""
    from io import BytesIO

    from openpyxl import load_workbook

    from backend.services.workpaper_summaries import system_refreshed
    from backend.services.workpaper_xlsx import render
    from backend.services.workpapers import load

    w, _ = walked
    snapshot = await load(w.x.runner.deps, w.x.eid)
    papers = {p.ref: p for p in snapshot.workpapers.values()}
    paper = papers["D"]
    table = next(t for s in paper.content.sections for t in s.tables if t.id == "lead_schedule")
    assert [r["account_code"] for r in table.rows] == ["200", "260", "270"]
    sales, _, interest = table.rows
    assert (
        sales["unaudited_cents"],
        sales["adjustments_cents"],
        sales["audited_cents"],
        sales["prior_cents"],
        sales["change_cents"],
    ) == (-3243100, 0, -3243100, -420000, -2823100)
    assert (
        interest["unaudited_cents"],
        interest["adjustments_cents"],
        interest["audited_cents"],
        interest["uncorrected_cents"],
    ) == (0, -4853, -4853, 0)
    assert all(ref in sales["source_ref"] for ref in ("D-1", "D-2", "D-3", "D-4"))
    period = next(s for s in paper.content.sections if s.id == "procedures").paragraphs[0]
    assert "3 months" in period.en and "12 months" in period.en and "not directly comparable" in period.en
    assert "not restated" in period.en
    cash = next(t for s in papers["C"].content.sections for t in s.tables if t.id == "lead_schedule")
    assert [r["account_code"] for r in cash.rows] == ["090", "091"]
    sheet = load_workbook(BytesIO(render(paper, "en"))).active
    rows = list(sheet.iter_rows())
    first = next(row[0].row for row in rows if row[0].value == "200")
    assert sheet.cell(first, 5).value == f"=C{first}+D{first}"
    assert sheet.cell(first, 7).value == f"=E{first}-F{first}"
    assert sheet.cell(first + 3, 3).value == f"=SUM(C{first}:C{first + 2})"
    assert any(cell.value == "✓" for row in rows for cell in row)
    for ref in ("C", "D"):
        spec = next(s for s in w.x.runner.deps.library.workpapers.system if s.ref == ref)
        assert system_refreshed(papers[ref], spec, snapshot.sources, papers) is None  # 没有导引表之间的版本循环.


@pytest.mark.external
@shared_loop
async def test_lead_unverified_adjustment_and_tieout_failure(walked: tuple[Walk, httpx.AsyncClient]) -> None:
    """有分录但未验证时不计已更正额; 未更正利息按收入贷方方向单列, 试算表异常时勾稽不冒充通过."""
    from dataclasses import replace

    from backend.contracts.models import MisstatementStatus
    from backend.services.workpaper_leads import lead_content
    from backend.services.workpapers import load

    w, _ = walked
    snapshot = await load(w.x.runner.deps, w.x.eid)
    sources = snapshot.sources
    paper = snapshot.workpapers["WP-D"]
    spec = next(s for s in w.x.runner.deps.library.workpapers.system if s.ref == "D")
    errors = [
        m.model_copy(
            update={
                "status": MisstatementStatus.uncorrected,
                "corrected_cents": 0,
                "residual_cents": None,
                "verified_by_judgment_id": None,
            }
        )
        if m.amount_cents == 4853
        else m
        for m in sources.misstatements
    ]
    unverified = replace(sources, misstatements=errors)
    content = lead_content(spec, unverified, paper, paper.content.header)
    table = next(t for s in content.sections for t in s.tables if t.id == "lead_schedule")
    interest = next(r for r in table.rows if r["account_code"] == "270")
    assert (interest["adjustments_cents"], interest["audited_cents"], interest["uncorrected_cents"]) == (0, 0, -4853)
    assert "C-4" in interest["source_ref"] and "X-" in interest["source_ref"]
    duplicate = next(r for r in sources.trial_balance if r.account_code == "200" and r.period == "current")
    malformed = replace(sources, trial_balance=[*sources.trial_balance, duplicate])
    content = lead_content(spec, malformed, paper, paper.content.header)
    tieout = next(t for s in content.sections for t in s.tables if t.id == "lead_tieout").rows[0]
    assert (tieout["status"], tieout["difference_cents"]) == ("✗", 3243100)
