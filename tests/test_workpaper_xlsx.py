"""从快照导出 Excel 的数据、公式、双语、版式与原文文本类型 (R331)."""

from io import BytesIO

import pytest
from openpyxl import load_workbook

from backend.contracts.models import WorkingLanguage
from backend.services.workpaper_xlsx import filename, money_format, render
from tests.workpaper_cases import d3_case

pytestmark = pytest.mark.external


@pytest.mark.parametrize("language", ["zh", "en"])
async def test_three_way_excel_cells_and_formulas(language: WorkingLanguage) -> None:
    _, _, paper = await d3_case()
    data = render(paper, language)
    sheet = load_workbook(BytesIO(data)).active
    assert sheet is not None
    assert filename(paper, "zh") == "D-3 凭证三单匹配 (REV-08).xlsx"
    assert "Demo Company (AU) Pty Ltd" in sheet.cell(2, 1).value
    assert "A6" in sheet.cell(4, 1).value
    assert ("待签署" if language == "zh" else "Awaiting sign-off") in sheet.cell(4, 9).value
    rows = list(sheet.iter_rows())
    first = next(row[0].row for row in rows if row[0].value == "ORC1024")
    assert [sheet.cell(first + n, 1).value for n in range(6)] == [
        "ORC1024",
        "ORC1020",
        "ORC1021",
        "ORC1022",
        "ORC1023",
        "ORC1026",
    ]
    assert sheet.cell(first + 5, 6).value == "XJ-580-2"
    assert (sheet.cell(first, 15).value, sheet.cell(first + 5, 15).value) == ("—", "✗")
    assert sheet.cell(first + 6, 5).value == f"=SUM(E{first}:E{first + 5})"
    assert sheet.cell(first + 6, 12).value == f"=SUM(L{first}:L{first + 5})"
    assert sum(sheet.cell(first + n, 5).value for n in range(6)) == 9900
    assert (
        sum(
            sheet.cell(first + n, 12).value
            for n in range(6)
            if isinstance(sheet.cell(first + n, 12).value, int | float)
        )
        == 6050
    )
    assert sheet.cell(first, 5).number_format == money_format("AUD")
    assert sheet.freeze_panes == "A6" and sheet.print_title_rows == "$1:$4"
    assert sheet.page_setup.orientation == "landscape" and sheet.page_setup.fitToWidth == 1
    assert sheet.cell(6, 1).fill.fgColor.rgb == "00E7EEF7"
    assert sheet.cell(first - 1, 1).fill.fgColor.rgb == "00F2F2F2"
    assert sheet.cell(first, 1).border.left.style == "thin" and sheet.cell(first, 1).alignment.wrap_text
    assert [tuple(cell.value for cell in row) for row in rows] == list(
        load_workbook(BytesIO(render(paper, language))).active.values
    )


async def test_original_text_is_not_an_excel_formula() -> None:
    _, _, paper = await d3_case()
    content = paper.content
    injected = '=HYPERLINK("https://invalid.example", "Synthetic original text")'
    content = content.model_copy(update={"header": content.header.model_copy(update={"client_name": injected})})
    # 测试一段裸原文, 不因前缀为等号而变成可执行公式.
    part = content.sections[-1].model_copy(update={"paragraphs": [injected]})
    content = content.model_copy(update={"sections": [*content.sections[:-1], part]})
    sheet = load_workbook(BytesIO(render(paper.model_copy(update={"content": content}), "en"))).active
    cell = next(cell for row in sheet for cell in row if cell.value == injected)
    assert cell.data_type == "s"
